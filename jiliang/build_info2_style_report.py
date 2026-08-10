#!/usr/bin/env python3
"""Create four info2-inspired, dataset-derived visuals and append them to the report."""
from __future__ import annotations

import csv
import math
import re
from collections import Counter
from itertools import combinations
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path('/Users/apple/Desktop/github.local/Li_Project/jiliang')
DATA = ROOT / 'output/01_merged_records.csv'
SOURCE = ROOT / 'output/report/04_False_Localizing_Sign_文献计量分析报告_extra图表增强版.docx'
ASSET_DIR = ROOT / 'output/figures/info2_style'
DEST = ROOT / 'output/report/05_False_Localizing_Sign_文献计量分析报告_info2图表增强版.docx'

COLORS = ['#e85d5d', '#47a85c', '#4d9ed7', '#d2af34', '#a46ccc', '#e58d47']
COUNTRY_ALIASES = {
    'UK': 'United Kingdom', 'U.K.': 'United Kingdom', 'England': 'United Kingdom',
    'USA': 'United States', 'U.S.A.': 'United States', 'United States of America': 'United States',
    'South Korea': 'Korea, Republic of', 'Republic of Korea': 'Korea, Republic of',
    'China, PR': 'China', 'Taiwan': 'Taiwan, Province of China',
}


def get_font(size: int, bold=False):
    candidates = [
        '/System/Library/Fonts/Supplemental/Arial Bold.ttf' if bold else '/System/Library/Fonts/Supplemental/Arial.ttf',
        '/System/Library/Fonts/Supplemental/Arial.ttf',
    ]
    for item in candidates:
        if Path(item).exists():
            return ImageFont.truetype(item, size)
    return ImageFont.load_default()


def read_records():
    with DATA.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def extract_countries(affiliations: str):
    text = (affiliations or '').lower()
    names = {
        'United States': ['united states', 'usa', 'u.s.a.'], 'United Kingdom': ['united kingdom', 'uk', 'england', 'scotland'],
        'China': ['china', 'beijing', 'shanghai'], 'Japan': ['japan'], 'India': ['india'], 'Canada': ['canada'],
        'Australia': ['australia'], 'Germany': ['germany'], 'France': ['france'], 'Italy': ['italy'], 'Spain': ['spain'],
        'Brazil': ['brazil'], 'Turkey': ['turkey'], 'Iran': ['iran'], 'Pakistan': ['pakistan'], 'Netherlands': ['netherlands'],
        'Switzerland': ['switzerland'], 'Sweden': ['sweden'], 'South Korea': ['south korea', 'republic of korea'],
        'Taiwan': ['taiwan'], 'Israel': ['israel'], 'Singapore': ['singapore'], 'Denmark': ['denmark'],
        'Belgium': ['belgium'], 'Austria': ['austria'], 'Greece': ['greece'], 'Mexico': ['mexico'],
    }
    return [country for country, terms in names.items() if any(re.search(r'(?<![a-z])' + re.escape(term) + r'(?![a-z])', text) for term in terms)]


def author_items(record):
    return [x.strip() for x in (record.get('authors') or '').split(';') if x.strip()]


def draw_network(path, title, counts, edges, kind):
    # Retain a small, legible collaboration core; edge means co-occurrence in one record.
    degree = Counter()
    for (a, b), weight in edges.items():
        degree[a] += weight; degree[b] += weight
    ranked = [x for x, _ in sorted(counts.items(), key=lambda kv: (-degree[kv[0]], -kv[1], kv[0])) if degree[x] > 0][:22]
    nodes = set(ranked)
    use_edges = [(a, b, w) for (a, b), w in edges.items() if a in nodes and b in nodes]
    W, H = 1800, 1120
    image = Image.new('RGB', (W, H), 'white'); draw = ImageDraw.Draw(image)
    title_font, label_font, small = get_font(35, True), get_font(19), get_font(17)
    draw.text((65, 32), title, font=title_font, fill='#111111')
    draw.line((65, 90, W - 65, 90), fill='#111111', width=3)
    cx, cy, rx, ry = W // 2, 550, 690, 330
    positions = {}
    for i, node in enumerate(ranked):
        angle = -math.pi / 2 + 2 * math.pi * i / max(len(ranked), 1)
        # Alternate two rings to avoid label collisions.
        scale = .76 if i % 2 else 1.0
        positions[node] = (cx + math.cos(angle) * rx * scale, cy + math.sin(angle) * ry * scale)
    for a, b, weight in use_edges:
        x1, y1 = positions[a]; x2, y2 = positions[b]
        draw.line((x1, y1, x2, y2), fill='#aeb9c3', width=2 + min(weight, 3), joint='curve')
    for i, node in enumerate(ranked):
        x, y = positions[node]; radius = 10 + int(7 * math.sqrt(counts[node]))
        fill = COLORS[i % len(COLORS)]
        draw.ellipse((x-radius, y-radius, x+radius, y+radius), fill=fill, outline='white', width=3)
        box = draw.textbbox((0, 0), node, font=label_font)
        draw.text((x-(box[2]-box[0])/2, y+radius+8), node, font=label_font, fill='#222222')
    note = ('Nodes are countries/regions; links indicate countries co-appearing in one record.' if kind == 'country'
            else 'Nodes are authors; links indicate co-authorship within one record. Node size reflects record count.')
    draw.text((65, H-70), note + ' Only the connected collaboration core is displayed.', font=small, fill='#333333')
    image.save(path)


def wrap(draw, text, font, width):
    words, out, line = text.split(), [], ''
    for word in words:
        test = (line + ' ' + word).strip()
        if line and draw.textlength(test, font=font) > width:
            out.append(line); line = word
        else: line = test
    return out + ([line] if line else [])


def draw_author_table(path, records, author_counts):
    top = author_counts.most_common(10); W, H = 1800, 1320
    im = Image.new('RGB', (W, H), '#fff9bd'); draw = ImageDraw.Draw(im)
    title, head, body, tiny = get_font(34, True), get_font(25, True), get_font(23), get_font(17)
    draw.text((60, 42), 'Table D1. Top 10 productive authors in false localizing sign research', font=title, fill='#111')
    draw.line((60, 108, W-60, 108), fill='#111', width=3)
    xs = [60, 420, 690]; labels = ['Authors', 'Documents', 'Representative publication in the present dataset']
    for x, lab in zip(xs, labels): draw.text((x, 128), lab, font=head, fill='#111')
    draw.line((60, 174, W-60, 174), fill='#111', width=2)
    y = 192
    for author, n in top:
        candidates = [r for r in records if author in author_items(r)]
        chosen = sorted(candidates, key=lambda r: int(r.get('year') or 0), reverse=True)[0] if candidates else {}
        article = re.sub(r'\s+', ' ', chosen.get('title') or '').strip()
        lines = wrap(draw, article, body, 1030)[:3]
        row_h = max(82, 25 + 30*len(lines))
        draw.text((60, y+12), author, font=body, fill='#111')
        draw.text((430, y+12), f'{n} ({n / len(records) * 100:.2f}%)', font=body, fill='#111')
        for j, line in enumerate(lines): draw.text((690, y+12+j*30), line, font=body, fill='#111')
        draw.line((60, y+row_h, W-60, y+row_h), fill='#777', width=1); y += row_h
    draw.text((60, H-60), 'Documents are counts within the 159 included records; the final column lists one recent publication for each author.', font=tiny, fill='#333')
    im.save(path)


def draw_journal_chart(path, records):
    counts = Counter(r.get('journal', '').strip() for r in records if r.get('journal', '').strip())
    top = counts.most_common(15); W, H = 1800, 1120
    im = Image.new('RGB', (W, H), 'white'); draw = ImageDraw.Draw(im)
    title, label, tick = get_font(31, True), get_font(21), get_font(19)
    draw.text((620, 30), 'Top 15 journals publishing on false localizing signs', font=title, fill='#333')
    left, right, top_y, bottom = 530, 1710, 105, 1010; max_n = max(n for _, n in top)
    for i, (name, n) in enumerate(top):
        y = top_y + i * 58; short = name if len(name) < 49 else name[:48] + '…'
        bbox = draw.textbbox((0, 0), short, font=label); draw.text((left-20-(bbox[2]-bbox[0]), y+8), short, font=label, fill='#333')
        length = (right-left) * n/max_n; draw.rectangle((left, y, left+length, y+38), fill='#3e789d')
    draw.line((left, bottom, right, bottom), fill='#555', width=2)
    for v in range(0, max_n+1, 1):
        x = left + (right-left)*v/max_n; draw.text((x-6, bottom+12), str(v), font=tick, fill='#444')
    draw.text((1010, 1070), 'Number of publications', font=label, fill='#333')
    im.save(path)


def set_font(run):
    # Songti SC is available to the LibreOffice renderer and preserves Chinese
    # appendix prose in the final PDF/Word rendering.
    run.font.name = 'Songti SC'; run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Songti SC')


def add_text(doc, text):
    p = doc.add_paragraph(); set_font(p.add_run(text)); return p


def add_heading(doc, text, level):
    p = doc.add_heading(level=level); set_font(p.add_run(text)); return p


def add_image(doc, path, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path), width=Inches(6.25))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = p.add_run(caption); run.italic = True; run.font.size = Pt(9); set_font(run)


def main():
    records = read_records(); ASSET_DIR.mkdir(parents=True, exist_ok=True)
    author_counts, author_edges, country_counts, country_edges = Counter(), Counter(), Counter(), Counter()
    for record in records:
        authors = list(dict.fromkeys(author_items(record))); author_counts.update(authors)
        author_edges.update(combinations(sorted(authors), 2))
        countries = list(dict.fromkeys(extract_countries(record.get('affiliations') or ''))); country_counts.update(countries)
        country_edges.update(combinations(sorted(countries), 2))
    assets = {
        'country': ASSET_DIR / 'fig_d1_country_collaboration_network.png',
        'authors_table': ASSET_DIR / 'table_d1_top_authors.png',
        'authors_network': ASSET_DIR / 'fig_d2_author_collaboration_network.png',
        'journals': ASSET_DIR / 'fig_d3_top_journals.png',
    }
    draw_network(assets['country'], 'Country/region collaboration network', country_counts, country_edges, 'country')
    draw_author_table(assets['authors_table'], records, author_counts)
    draw_network(assets['authors_network'], 'Author co-authorship network', author_counts, author_edges, 'author')
    draw_journal_chart(assets['journals'], records)
    doc = Document(SOURCE); doc.add_page_break()
    add_heading(doc, '附录 D  参照 info2 示例形式的合作网络与核心作者分析', 1)
    add_text(doc, '本附录参照 info2 中的国家合作网络、作者合作网络、核心作者表和期刊条形图的版式，使用本研究纳入的 159 篇假定位体征文献重新绘制。网络连线均仅表示同一篇题录中的共同署名或共同出现，不代表引文关系、机构协作强度或因果关系。')
    add_heading(doc, 'D.1 国家/地区合作结构', 2)
    add_text(doc, '根据题录地址字段识别国家/地区，并在同一篇记录出现多个国家/地区时建立合作连线。由于部分早期记录或题录缺少地址信息，该图用于描述可识别的合作结构。')
    add_image(doc, assets['country'], 'Fig. D1. Co-authorship network of countries/regions in false localizing sign research.')
    doc.add_page_break()
    add_heading(doc, 'D.2 核心作者与作者合作网络', 2)
    add_text(doc, '表 D1 展示按本研究数据集发文量排序的前 10 位作者；图 D2 进一步展示存在共同署名关系的作者合作核心。')
    add_image(doc, assets['authors_table'], 'Table D1. Top 10 productive authors in false localizing sign research.')
    doc.add_page_break()
    add_image(doc, assets['authors_network'], 'Fig. D2. Co-authorship network of authors in false localizing sign research.')
    doc.add_page_break()
    add_heading(doc, 'D.3 核心期刊', 2)
    add_text(doc, '图 D3 沿用示例的横向条形图形式，展示本数据集的前 15 位发文期刊。')
    add_image(doc, assets['journals'], 'Fig. D3. Top 15 journals publishing on false localizing signs.')
    doc.save(DEST); print(DEST)


if __name__ == '__main__': main()
