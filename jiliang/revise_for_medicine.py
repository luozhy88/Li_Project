#!/usr/bin/env python3
"""Revise report v06 into v07 (Medicine submission-oriented revision).

Fixes applied (acting as reviewer / bioinformatics expert):
 1. Correct retrieval/dedup numbers (125+89 -> 159, removed 55; not "123+36 as retrieved").
 2. Remove placeholders ("如已导入"), outdated generation date.
 3. Author stats use canonicalized names (parser fix upstream); text updated to new counts.
 4. Consistent figure/table numbering: body Fig. 1-7, appendices A/B/C.
 5. Rebuild appendix A record list without 'nan' fields.
 6. Regenerate first-author table image labelled Table C1 and swap it in.
 7. Soften "first bibliometric analysis" claim; update limitations (WoS now included).
 8. Chinese typography: space between CJK characters and digits.
"""
from __future__ import annotations

import csv
import importlib.util
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

ROOT = Path('/Users/apple/Desktop/github.local/Li_Project/jiliang')
SOURCE = ROOT / 'output/report/06_False_Localizing_Sign_文献计量分析报告_图表修订版.docx'
DEST = ROOT / 'output/report/07_False_Localizing_Sign_文献计量分析报告_Medicine投稿修订版.docx'
DATA = ROOT / 'output/01_merged_records.csv'
TABLE_C1 = ROOT / 'output/figures/geo_style/table_c1_top_first_authors.png'


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def set_font(run):
    run.font.name = 'Songti SC'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Songti SC')


def rewrite(p, text):
    """Replace a paragraph's text, keeping its first run's formatting."""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)


# ---------------------------------------------------------------- token renumbering
TOKEN_MAP = {
    # body figures: types figure becomes Fig. 2, others shift by one
    'Fig.3-1': 'Fig. 2', 'Fig.2': 'Fig. 3', 'Fig.3': 'Fig. 4', 'Fig.4': 'Fig. 5',
    'Fig.5': 'Fig. 6', 'Fig.6': 'Fig. 7',
    '图2': '图 3', '图3': '图 4', '图4': '图 5', '图5': '图 6', '图6': '图 7',
    # appendix C (extra-style) -> appendix B
    'Fig.C1': 'Fig. B1', 'Fig.C2': 'Fig. B2', 'Fig.C3': 'Fig. B3',
    'TableC1': 'Table B1', 'TableC2': 'Table B2',
    'C.1': 'B.1', 'C.2': 'B.2',
    # appendix D (geo-style) -> appendix C
    'Fig.D1': 'Fig. C1', 'Fig.D2': 'Fig. C2', 'Fig.D3': 'Fig. C3',
    'Fig.D4': 'Fig. C4', 'Fig.D5': 'Fig. C5', 'TableD1': 'Table C1',
    '图D1': '图 C1', '图D2': '图 C2', '图D3': '图 C3', '图D4': '图 C4', '图D5': '图 C5',
    '表D1': '表 C1',
    'D.1': 'C.1', 'D.2': 'C.2', 'D.3': 'C.3', 'D.4': 'C.4',
}
TOKEN_RE = re.compile(
    r'(Fig\.\s*\d+-\d+|Fig\.\s*[A-Z]\d+|Fig\.\s*\d+|Table\s+[A-Z]\d+|'
    r'图\s*[A-Z]\d+|图\s*\d+|表\s*[A-Z]\d+|\b[CD]\.\d+)')


def renumber(text):
    return TOKEN_RE.sub(lambda m: TOKEN_MAP.get(re.sub(r'\s+', '', m.group(0)), m.group(0)), text)


# ---------------------------------------------------------------- specific rewrites
ABSTRACT = ('目的：对假定位体征（false localizing sign）相关文献进行文献计量学分析，梳理该领域的研究现状、'
            '热点与发展趋势。方法：在 PubMed 中以 Title/Abstract 字段、在 Web of Science（WoS）核心合集中以 '
            'Topic 字段检索 "false localizing sign"、"false localising sign" 与 "pseudo-localization sign" '
            '等拼写变体；对两库题录进行解析并按 PMID、DOI 与归一化标题去重后，采用 Python 进行描述性统计、'
            '文献体裁构成、期刊/国家/地区/作者分布及关键词共现分析。结果：共纳入 159 篇文献（PubMed 123 篇、'
            'WoS 36 篇），时间跨度为 1904—2026 年；文献体裁以病例报告为主（72 篇，45.3%），综述（含系统综述）'
            '11 篇（6.9%）；发文量总体呈波动上升趋势，刊载期刊集中于神经病学、神经外科学与神经眼科学领域；'
            '高频关键词包括 Magnetic Resonance Imaging、Diagnosis, Differential、Brain Neoplasms、'
            'Tomography, X-Ray Computed、Cervical Vertebrae 等。结论：假定位体征的研究仍以病例报告与影像学'
            '特征为主，未来可进一步开展多中心、大样本的临床与影像学研究。')

S31 = ('经检索，PubMed 获得 125 条记录，WoS 获得 89 条记录；依次按 PMID、DOI 与归一化标题去重，'
       '共移除 55 条重复记录，最终纳入 159 篇文献进行分析（PubMed 来源 123 篇，WoS 来源 36 篇），'
       '时间跨度为 1904—2026 年。其中具有 DOI 的记录 135 条，具有摘要的记录 123 条。'
       '文献体裁构成方面（按优先级互斥归类）：病例报告 72 篇（45.3%），综述（含系统综述）11 篇（6.9%），'
       '临床试验 1 篇（0.6%），来信/述评/评论 4 篇（2.5%），历史/传记类 4 篇（2.5%），'
       '其他期刊论文 66 篇（41.5%），其他类型 1 篇（0.6%），见图 2。')

S34 = ('基于题录中全部作者的地址字段提取国家/地区（同一国家/地区在同一篇文献中仅计一次），'
       '发文量前 3 位为：美国（24 篇）、印度（13 篇）、日本（12 篇）。图 4 展示了发文量前 10 位的'
       '国家/地区；附录 C 进一步给出了世界分布地图（图 C1）、前 10 位国家/地区条形图（图 C2）与'
       '国际合作网络（图 C3）。')

S35 = ('对作者姓名拼写变体进行规范化合并后，发文量最多的作者为 Carrasco-Moro R 与 Pascual JM'
       '（各 5 篇），其次为 Martínez-San Millán JS、Schievink WI 与 Maya MM（各 4 篇）。'
       '图 5 展示了前 10 位核心作者；附录 C 表 C1 另按第一作者口径统计（每篇文献仅计入其第一作者一次），'
       '以避免同一篇文献在多位作者之间重复计数。')

DISCUSSION = ('据我们所知，本研究是首次对假定位体征相关文献进行的系统计量分析。结果显示，该领域文献数量'
              '总体呈增长态势，尤其近十年病例报告与影像学研究增多，反映了神经影像技术对识别罕见假定位体征的'
              '促进作用。从文献体裁看，病例报告占 45.3%，提示该领域证据仍以个案经验积累为主，缺乏大样本队列'
              '与对照研究。从期刊分布看，研究主要发表于神经病学、神经外科学与神经眼科学期刊，提示该主题具有'
              '鲜明的专科交叉特征。国家/地区分布上，美国、印度、日本为该领域的主要贡献国家，但国际合作连线'
              '仍较稀疏，提示跨国协作尚不充分。高频关键词集中于磁共振成像、颅内肿瘤、颅内压异常及颅神经麻痹'
              '等，说明假定位体征的研究热点与神经影像、颅内压管理及颅神经病变密切相关。\n'
              '本研究的局限性包括：仅纳入 PubMed 与 WoS 两个数据库，未覆盖 Embase、Scopus 等来源，且未进行'
              '引文与共被引分析；机构与国家提取基于地址文本解析，可能存在少量误判；作者姓名虽经拼写变体'
              '规范化合并，仍可能存在同名异人的少量误判；关键词共现基于 MeSH 与作者关键词，未进行语义聚类。'
              '未来研究可结合 CiteSpace 的突现检测与 VOSviewer 的共被引分析，进一步揭示学科演化路径与'
              '知识基础。')


def main():
    # ---- regenerate the first-author table image labelled Table C1
    geo = load_module('geo', ROOT / 'build_geo_style_report.py')
    records = list(csv.DictReader(open(DATA, encoding='utf-8-sig', newline='')))
    geo.draw_first_author_table(TABLE_C1, records, label='Table C1')

    gen = load_module('gen', ROOT / '4_generate_report.py')

    doc = Document(SOURCE)

    # ---- pass 1: token renumbering on every paragraph
    for p in doc.paragraphs:
        new = renumber(p.text)
        if new != p.text:
            rewrite(p, new)

    # ---- pass 2: specific content rewrites
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('生成日期：'):
            rewrite(p, '生成日期：2026-08-10')
        elif t.startswith('目的：'):
            rewrite(p, ABSTRACT)
        elif t.startswith('关键词：假定位体征'):
            rewrite(p, '关键词：假定位体征；文献计量学；PubMed；Web of Science；共现网络分析')
        elif '检索词覆盖' in t and '检索日期' not in t:
            rewrite(p, t + ' PubMed 检索日期为 2026 年 7 月 12 日，WoS 检索日期为 2026 年 7 月 19 日。')
        elif t.startswith('纳入标准：'):
            rewrite(p, '纳入标准：检索所得且题录信息完整的文献记录，文献体裁不限（期刊论文、病例报告、综述、'
                       '来信等均可纳入）；排除标准：重复发表的记录，以及标题、作者等关键字段缺失而无法解析的记录。')
        elif t.startswith('采用描述性统计'):
            rewrite(p, t + '作者姓名按连字符、空格、逗号与大小写差异进行拼写变体规范化合并；'
                           '世界分布图基于 Natural Earth 地图数据绘制。')
        elif '具有摘要的记录' in t:
            rewrite(p, S31)
        elif '按第一/通讯作者机构地址提取' in t:
            rewrite(p, S34)
        elif t.startswith('发文量最多的作者包括'):
            rewrite(p, S35)
        elif '本研究首次对假定位体征' in t:
            rewrite(p, DISCUSSION)
        elif t.startswith('附录：纳入文献清单'):
            rewrite(p, '附录 A  纳入文献清单（前 30 条）')
        elif t.startswith('附录 C  基于 extra'):
            rewrite(p, '附录 B  基于 extra 示例形式的补充分析')
        elif t.startswith('附录 D'):
            rewrite(p, '附录 C  国家/地区分布、合作网络与核心作者分析')

    # ---- pass 3: caption label punctuation ("Fig. 3 Top" -> "Fig. 3. Top")
    for p in doc.paragraphs:
        t = p.text
        if re.match(r'^(Fig\.|Table)\s', t.strip()):
            new = re.sub(r'^((?:Fig\.|Table)\s+[A-Z]?\d+)[^\d.]', r'\1. ', t.strip())
            new = re.sub(r'^((?:Fig\.|Table)\s+[A-Z]?\d+)\.\.$', r'\1.', new)
            if new != t.strip():
                rewrite(p, new)

    # ---- pass 4: rebuild appendix A record list without 'nan'
    paras = doc.paragraphs
    head = next(p for p in paras if p.text.strip().startswith('附录 A'))
    intro = ('以下列出纳入分析的 159 条记录中的前 30 条，完整清单见 output/01_merged_records.csv。')
    entries = []
    for i, row in enumerate(records[:30], 1):
        entries.append(f'{i}. {gen.format_reference(row)}')
    # remove everything between appendix A heading and appendix B heading
    el = head._p
    for sib in list(el.itersiblings()):
        if sib.tag == qn('w:p'):
            txt = ''.join(sib.itertext())
            if txt.strip().startswith('附录 B'):
                break
        el.getparent().remove(sib)
    anchor = head._p
    for text in [intro] + entries:
        p = doc.add_paragraph()
        set_font(p.add_run(text))
        anchor.addnext(p._p)
        anchor = p._p

    # ---- pass 5: swap the first-author table image with the Table C1 version
    for p in doc.paragraphs:
        if p.text.strip().startswith('Table C1.'):
            el = p._p.getprevious()
            while el is not None and not el.findall('.//' + qn('w:drawing')):
                el = el.getprevious()
            assert el is not None, '未找到 Table C1 对应的图片段落'
            from docx.text.paragraph import Paragraph
            pic_par = Paragraph(el, p._parent)
            pic_par.clear()
            pic_par.add_run().add_picture(str(TABLE_C1), width=Inches(6.25))
            break

    # ---- pass 6: CJK-digit spacing for Chinese paragraphs
    for p in doc.paragraphs:
        t = p.text
        if not re.search(r'[一-鿿]', t):
            continue
        new = re.sub(r'(?<=[一-鿿])(?=\d)', ' ', t)
        new = re.sub(r'(?<=\d)(?=[一-鿿])', ' ', new)
        if new != t:
            rewrite(p, new)

    doc.save(DEST)
    print(DEST)


if __name__ == '__main__':
    main()
