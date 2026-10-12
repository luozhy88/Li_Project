#!/usr/bin/env python3
"""Synchronize Table B2 with the data-driven Fig. B2 clusters."""
from __future__ import annotations

import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches


ROOT = Path('/Users/apple/Desktop/github.local/Li_Project/jiliang')
STATS = ROOT / 'output/logs/b1_own_stats.json'
SOURCE = ROOT / 'output/report/10_False_Localizing_Sign_文献计量分析报告_Medicine投稿修订版v10.docx'
DEST = ROOT / 'output/report/10_False_Localizing_Sign_文献计量分析报告_Medicine投稿修订版v10.docx'
ASSET = ROOT / 'input/info/extra_style_assets/table_b2_keyword_clusters_synced.png'


def get_font(size: int, bold: bool = False):
    path = '/System/Library/Fonts/Supplemental/Arial Bold.ttf' if bold else '/System/Library/Fonts/Supplemental/Arial.ttf'
    return ImageFont.truetype(path, size)


def wrap(draw, text, font, max_width):
    words = str(text).split()
    lines, current = [], ''
    for word in words:
        proposal = (current + ' ' + word).strip()
        if current and draw.textlength(proposal, font=font) > max_width:
            lines.append(current)
            current = word
        else:
            current = proposal
    if current:
        lines.append(current)
    return lines or ['']


def make_table():
    stats = json.loads(STATS.read_text(encoding='utf-8'))
    clustering = stats['clustering']
    clusters = clustering['clusters']

    width = 2400
    title_font = get_font(40, True)
    header_font = get_font(28, True)
    body_font = get_font(25)
    small_font = get_font(21)
    col_widths = [190, 170, 560, 1380]
    left = 50
    top = 40
    header_h = 72
    row_h = 150
    title_h = 85
    note_h = 100
    height = top + title_h + header_h + row_h * len(clusters) + note_h

    # Pure white canvas and white cells, as requested.
    image = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(image)
    draw.text((left, top), 'Table B2. Data-driven keyword clusters corresponding to Fig. B2', fill='black', font=title_font)

    x_positions = [left]
    for w in col_widths:
        x_positions.append(x_positions[-1] + w)
    y = top + title_h
    headers = ['Cluster', 'Nodes', 'LLR label', 'Representative keywords (top members)']
    for i, header in enumerate(headers):
        draw.rectangle((x_positions[i], y, x_positions[i + 1], y + header_h), fill='white', outline='black', width=2)
        draw.text((x_positions[i] + 14, y + 20), header, fill='black', font=header_font)

    y += header_h
    for cluster in clusters:
        row = [
            f"#{cluster['id']}",
            str(cluster['size']),
            cluster['label'],
            '; '.join(cluster['members_top']),
        ]
        for i, value in enumerate(row):
            draw.rectangle((x_positions[i], y, x_positions[i + 1], y + row_h), fill='white', outline='black', width=1)
            font = body_font if i != 3 else small_font
            lines = wrap(draw, value, font, col_widths[i] - 28)
            line_h = 31 if i != 3 else 27
            total_h = min(len(lines), 4) * line_h
            yy = y + max(14, (row_h - total_h) // 2)
            for line in lines[:4]:
                draw.text((x_positions[i] + 14, yy), line, fill='black', font=font)
                yy += line_h
        y += row_h

    note = (f"Note: same 159 records, 70-node/206-link pruned network, greedy-modularity clustering, "
            f"LLR labels; Q = {clustering['modularity_Q']:.3f}, weighted mean S = {clustering['weighted_mean_silhouette_S']:.3f}.")
    draw.text((left, y + 24), note, fill='black', font=small_font)
    ASSET.parent.mkdir(parents=True, exist_ok=True)
    image.save(ASSET)


def replace_table_image():
    doc = Document(SOURCE)
    paragraphs = doc.paragraphs
    target = next(p for p in paragraphs if p.text.strip().startswith('Table B2.'))
    target.clear()
    target.add_run('Table B2. Data-driven keyword clusters corresponding to Fig. B2.')
    element = target._p.getprevious()
    while element is not None and not element.findall('.//' + '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'):
        element = element.getprevious()
    if element is None:
        raise RuntimeError('Could not locate the image paragraph preceding Table B2.')

    from docx.text.paragraph import Paragraph
    image_paragraph = Paragraph(element, target._parent)
    image_paragraph.clear()
    image_paragraph.alignment = target.alignment
    image_paragraph.add_run().add_picture(str(ASSET), width=Inches(6.25))
    doc.save(DEST)


if __name__ == '__main__':
    make_table()
    replace_table_image()
    print(DEST)
