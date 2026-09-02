#!/usr/bin/env python3
"""Update report v09 -> v10: swap in the CiteSpace-style annual-ring Fig. C3
and update its body text and caption.

Nodes are now drawn as concentric rings (one ring per 5-year publication
period; ring color encodes the period on a shared 1975-2025 scale; ring area
is proportional to the publication count in that period).
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

ROOT = Path('/Users/apple/Desktop/github.local/Li_Project/jiliang')
SOURCE = ROOT / 'output/report/09_False_Localizing_Sign_文献计量分析报告_Medicine投稿修订版v3.docx'
DEST = ROOT / 'output/report/10_False_Localizing_Sign_文献计量分析报告_Medicine投稿修订版v4.docx'
IMG = ROOT / 'output/figures/geo_style/fig_e2_country_network_glow.png'

BODY = ('图 C3 展示国家/地区间的合作网络：节点面积与总发文量严格成正比（见图内 Legend 示例）；'
        '节点呈年轮状同心环结构，每一环代表一个有发文记录的 5 年时间切片，'
        '环颜色对应发文年代（见色条，1975—2025 年），环面积与该时段发文量成正比；'
        '红色连线表示两国/地区在同一篇文献中共同署名；'
        '无合作连线的国家/地区（孤立节点）统一排列于图谱底部区域，'
        '中央画幅保留给具有合作关系的国家/地区聚类。')

CAPTION = ('Fig. C3. Network visualization of international collaboration among countries in '
           'false localizing sign research. Each node represents a country, with node area '
           'proportional to total publication output. Concentric rings represent 5-year periods '
           'with publication activity: ring color indicates the publication period (1975-2025, '
           'see color scale) and ring area is proportional to the number of publications in that '
           'period. Links indicate co-authorship within one record; countries without '
           'co-authorship links (isolated nodes) are arranged in the bottom strip.')


def rewrite(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)


def main():
    doc = Document(SOURCE)
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('图 C3 展示国家/地区间的合作网络'):
            rewrite(p, BODY)
            print('正文 C.2 段落已更新')
        elif t.startswith('Fig. C3.'):
            el = p._p.getprevious()
            while el is not None and not el.findall('.//' + qn('w:drawing')):
                el = el.getprevious()
            assert el is not None, '未找到 Fig. C3 对应的图片段落'
            pic_par = Paragraph(el, p._parent)
            pic_par.clear()
            pic_par.alignment = p.alignment  # 保持居中
            pic_par.add_run().add_picture(str(IMG), width=Inches(6.25))
            rewrite(p, CAPTION)
            print('Fig. C3 图片与图注已更新')
    doc.save(DEST)
    print(DEST)


if __name__ == '__main__':
    main()
