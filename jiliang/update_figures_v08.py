#!/usr/bin/env python3
"""Update report v07 -> v08: swap in revised figures (reviewer figure-quality fixes).

Replaced images (regenerated from own data after thesaurus merging, white
background, modularity clustering and label-overlap prevention):
  Fig. 6  <- output/figures/fig5_keyword_network.png        (keyword network)
  Fig. 7  <- output/figures/fig6_keyword_wordcloud.png      (word cloud)
  Fig. B1 <- input/info/.../fig10_keyword_network_and_clusters.png
  Fig. B2 <- input/info/.../fig11_topic_time_distribution.png
Captions updated accordingly (Fig. B2 notes chronological arrangement and the
n >= 5 filter).
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

ROOT = Path('/Users/apple/Desktop/github.local/Li_Project/jiliang')
SOURCE = ROOT / 'output/report/07_False_Localizing_Sign_文献计量分析报告_Medicine投稿修订版.docx'
DEST = ROOT / 'output/report/08_False_Localizing_Sign_文献计量分析报告_Medicine投稿修订版v2.docx'

REPLACEMENTS = {
    'Fig. 6.': (
        ROOT / 'output/figures/fig5_keyword_network.png',
        'Fig. 6. Keyword co-occurrence network based on MeSH terms and author keywords. '
        'Synonymous variants were merged using a thesaurus file before analysis; node colors '
        'indicate modularity-based clusters and node size is proportional to keyword frequency.'),
    'Fig. 7.': (
        ROOT / 'output/figures/fig6_keyword_wordcloud.png',
        None),  # caption unchanged
    'Fig. B1.': (
        ROOT / 'input/info/false_localizing_info_style_figures/fig10_keyword_network_and_clusters.png',
        'Fig. B1. Keyword co-occurrence network (A) and modularity-based clustered thematic '
        'structure (B). Synonymous variants were merged using a thesaurus file before analysis.'),
    'Fig. B2.': (
        ROOT / 'input/info/false_localizing_info_style_figures/fig11_topic_time_distribution.png',
        'Fig. B2. Observed time span of major research topics. Topics arranged chronologically '
        'by initial appearance year; only topics with n >= 5 records are shown.'),
}


def rewrite(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)


def main():
    doc = Document(SOURCE)
    for p in doc.paragraphs:
        t = p.text.strip()
        for prefix, (img, caption) in REPLACEMENTS.items():
            if not t.startswith(prefix):
                continue
            el = p._p.getprevious()
            while el is not None and not el.findall('.//' + qn('w:drawing')):
                el = el.getprevious()
            assert el is not None, f'未找到 {prefix} 对应的图片段落'
            pic_par = Paragraph(el, p._parent)
            pic_par.clear()
            pic_par.add_run().add_picture(str(img), width=Inches(6.25))
            if caption:
                rewrite(p, caption)
            print(f'已替换 {prefix} 图片 -> {img.name}')
            break
    doc.save(DEST)
    print(DEST)


if __name__ == '__main__':
    main()
