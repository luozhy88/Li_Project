#!/usr/bin/env python3
"""Create evidence-based bibliometric visuals in the style of input/info.

The source data do not contain citation counts or cited-reference lists.  This
script therefore uses only fields available in 01_merged_records.csv and labels
every output accordingly.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path
import importlib.util
import re

import yaml

import matplotlib.pyplot as plt
from matplotlib import gridspec
from matplotlib.patches import FancyBboxPatch
import networkx as nx
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

ROOT = Path('/Users/apple/Desktop/github.local/Li_Project/jiliang')
DATA = ROOT / 'output/01_merged_records.csv'
TABLES = ROOT / 'output/tables'
SOURCE_DOCX = ROOT / 'output/report/False_Localizing_Sign_文献计量分析报告.docx'
WORK = Path('/Users/apple/Desktop/github.local/Li_Project/jiliang/input/info/false_localizing_info_style_figures')
DEST = Path('/Users/apple/Desktop/github.local/Li_Project/jiliang/input/info/False_Localizing_Sign_文献计量分析报告_数据图表增强版.docx')

YELLOW = '#fff6ad'
HEADER = '#f7e882'
INK = '#141414'
RED = '#e74c3c'
BLUE = '#76c8cf'
PALETTE = ['#e76f51', '#2a9d8f', '#457b9d', '#e9c46a', '#9b5de5', '#f4a261']


def style():
    plt.rcParams.update({
        'font.family': 'Arial', 'font.size': 9, 'axes.edgecolor': INK,
        'text.color': INK, 'figure.facecolor': YELLOW, 'axes.facecolor': YELLOW,
        'savefig.facecolor': YELLOW,
    })


def save(fig, name, facecolor=YELLOW):
    path = WORK / name
    fig.savefig(path, dpi=240, bbox_inches='tight', facecolor=facecolor)
    plt.close(fig)
    return path


# ---- 同义词归一化（复用 3_keyword_analysis.py 的实现与 config.yaml 同义词表）----
def _kw_module():
    spec = importlib.util.spec_from_file_location('kw_analysis', ROOT / '3_keyword_analysis.py')
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


_KW = _kw_module()
_CONFIG = yaml.safe_load(open(ROOT / 'config.yaml', encoding='utf-8'))
_STOPWORDS = {w.lower() for w in _CONFIG.get('keyword', {}).get('stopwords', [])}
_MAPPER_CACHE = None


def kw_mapper(df):
    global _MAPPER_CACHE
    if _MAPPER_CACHE is None:
        _MAPPER_CACHE = _KW.build_keyword_mapper(df, _STOPWORDS, _CONFIG)
    return _MAPPER_CACHE


def map_terms(items, mapper):
    return list(dict.fromkeys(mapper.get(x, x) for x in items))


def keywords(value):
    if pd.isna(value):
        return []
    return [x.strip().rstrip('*').strip() for x in str(value).split(';') if x.strip()]


def make_representative_records(df):
    # "Recent representative records" rather than a high-citation claim.
    records = df.dropna(subset=['title', 'year']).copy()
    records = records[records['title'].str.len() > 8]
    records = records.sort_values(['year', 'title'], ascending=[False, True]).head(10)
    fig, ax = plt.subplots(figsize=(12.2, 7.1))
    ax.set_axis_off()
    ax.text(0.01, .965, 'Table 1. Ten recent representative records on false localizing signs', fontsize=15, weight='bold', va='top')
    ax.plot([.01, .99], [.92, .92], color=INK, lw=1.6)
    columns = [('Year', .01), ('Title', .10), ('Journal', .61), ('DOI', .82)]
    for label, x in columns:
        ax.text(x, .885, label, weight='bold', fontsize=10, va='center')
    ax.plot([.01, .99], [.855, .855], color=INK, lw=1.0)
    y = .82
    for _, r in records.iterrows():
        title = re.sub(r'\s+', ' ', str(r['title'])).strip()
        journal = re.sub(r'\s+', ' ', str(r['journal'] if pd.notna(r['journal']) else '')).strip()
        doi = str(r['doi'] if pd.notna(r['doi']) else '')
        def short(text, n): return text if len(text) <= n else text[:n - 1] + '…'
        ax.text(.01, y, str(int(r.year)), fontsize=8, va='top')
        ax.text(.10, y, short(title, 70), fontsize=8, va='top', wrap=True)
        ax.text(.61, y, short(journal, 28), fontsize=8, va='top', wrap=True)
        ax.text(.82, y, short(doi, 25), fontsize=8, va='top', wrap=True)
        y -= .072
    ax.plot([.01, .99], [.09, .09], color=INK, lw=1.0)
    ax.text(.01, .04, 'Note: ranked by recency for transparent data presentation; citation counts were not available in the source records.', fontsize=8)
    return save(fig, 'fig7_recent_representative_records.png')


def make_journal_table():
    d = pd.read_csv(TABLES / '02_top_journals.csv').head(10).copy()
    total = 159
    fig, ax = plt.subplots(figsize=(12.2, 6.2)); ax.set_axis_off()
    ax.text(.01, .965, 'Table 2. Top 10 journals in false localizing sign research', fontsize=15, weight='bold', va='top')
    ax.plot([.01, .99], [.92, .92], color=INK, lw=1.6)
    headers = [('Journal', .01), ('Articles count', .64), ('Share of records', .81)]
    for label, x in headers: ax.text(x, .885, label, weight='bold', fontsize=10)
    ax.plot([.01, .99], [.855, .855], color=INK, lw=1.0)
    y = .815
    for _, r in d.iterrows():
        journal = str(r['Journal']); journal = journal if len(journal) < 58 else journal[:57] + '…'
        ax.text(.01, y, journal, fontsize=10, va='center')
        ax.text(.68, y, str(int(r['Publications'])), fontsize=10, va='center', ha='center')
        ax.text(.86, y, f"{100*r['Publications']/total:.1f}%", fontsize=10, va='center', ha='center')
        y -= .069
    ax.plot([.01, .99], [.105, .105], color=INK, lw=1.0)
    ax.text(.01, .045, 'Share is calculated using all 159 included records. Journal country and impact factor are intentionally not shown because they are not contained in the dataset.', fontsize=8)
    return save(fig, 'fig8_top_journals_table.png')


def make_author_table(df):
    top = pd.read_csv(TABLES / '02_top_authors.csv').head(10)
    fig, ax = plt.subplots(figsize=(12.2, 6.5)); ax.set_axis_off()
    ax.text(.01, .965, 'Table 3. Top 10 productive authors in false localizing sign research', fontsize=15, weight='bold', va='top')
    ax.plot([.01, .99], [.92, .92], color=INK, lw=1.6)
    for label, x in [('Author', .01), ('Documents', .23), ('Representative publication in the present dataset', .39)]:
        ax.text(x, .885, label, weight='bold', fontsize=10)
    ax.plot([.01, .99], [.855, .855], color=INK, lw=1.0)
    y = .815
    for _, r in top.iterrows():
        author = r['Author']
        candidates = df[df['authors'].fillna('').str.contains(re.escape(author), regex=True)]
        title = '' if candidates.empty else str(candidates.sort_values('year', ascending=False).iloc[0]['title'])
        title = re.sub(r'\s+', ' ', title).strip()
        title = title if len(title) < 78 else title[:77] + '…'
        ax.text(.01, y, author, fontsize=10, va='center')
        ax.text(.28, y, f"{int(r['Publications'])} ({100*r['Publications']/159:.2f}%)", fontsize=10, va='center', ha='center')
        ax.text(.39, y, title, fontsize=8.5, va='center', wrap=True)
        y -= .069
    ax.plot([.01, .99], [.105, .105], color=INK, lw=1.0)
    ax.text(.01, .045, 'Documents are counts within the 159 included records; the final column gives one record authored by each listed researcher.', fontsize=8)
    return save(fig, 'fig9_top_authors_table.png')


def network_data(df):
    counts = Counter()
    edges = Counter()
    mapper = kw_mapper(df)
    for val in df['keywords']:
        items = map_terms(keywords(val), mapper)
        counts.update(items)
    # Remove demographic/indexing terms and retain a legible set of substantive
    # research concepts.  This avoids a decorative but unreadable dense graph.
    generic = {
        'humans', 'male', 'female', 'middle aged', 'adult', 'aged', 'child',
        'young adult', 'adolescent', 'aged, 80 and over', 'retrospective studies',
        'history, 19th century', 'history, 20th century', 'case reports',
    }
    candidates = [(x, n) for x, n in counts.items()
                  if n >= 5 and x.lower() not in generic
                  and not x.lower().startswith('false localiz')]
    selected = {x for x, _ in sorted(candidates, key=lambda item: (-item[1], item[0]))[:26]}
    for val in df['keywords']:
        items = sorted(set(x for x in map_terms(keywords(val), mapper) if x in selected))
        edges.update(combinations(items, 2))
    G = nx.Graph()
    for item in selected: G.add_node(item, frequency=counts[item])
    for (a, b), weight in edges.items():
        if weight >= 2: G.add_edge(a, b, weight=weight)
    return G, counts


def make_network_and_clusters(df):
    G, _ = network_data(df)
    pos = nx.spring_layout(G, seed=19, k=1.5, iterations=250)
    communities = list(nx.algorithms.community.greedy_modularity_communities(G)) if G.number_of_edges() else [set(G.nodes())]
    community_index = {n: i for i, c in enumerate(communities) for n in c}
    fig = plt.figure(figsize=(13, 6.8)); gs = gridspec.GridSpec(1, 2, width_ratios=[1.08, 1])
    for idx, title in enumerate(['(A) Keyword co-occurrence network', '(B) Clustered keyword network']):
        ax = fig.add_subplot(gs[idx]); ax.set_title(title, fontsize=12, weight='bold', loc='left')
        sizes = [100 + 65 * G.nodes[n]['frequency'] for n in G.nodes]
        widths = [.35 + .9 * data['weight'] for _, _, data in G.edges(data=True)]
        colors = '#8ecae6' if idx == 0 else [PALETTE[community_index[n] % len(PALETTE)] for n in G.nodes]
        nx.draw_networkx_edges(G, pos, width=widths, alpha=.35, edge_color='#7a7a7a', ax=ax)
        nx.draw_networkx_nodes(G, pos, node_size=sizes, node_color=colors, alpha=.9, edgecolors='white', linewidths=.8, ax=ax)
        ax.margins(0.22)  # 必须先固定坐标范围，再进行像素级标签防重叠
        texts = nx.draw_networkx_labels(G, pos, font_size=7, ax=ax)
        _KW.adjust_labels(ax, list(texts.values()))  # 防标签重叠
        ax.axis('off')
    fig.suptitle('Keyword analysis of false localizing sign research', fontsize=16, weight='bold', y=1.02)
    return save(fig, 'fig10_keyword_network_and_clusters.png', facecolor='white')


def make_timeline(df):
    # 数据驱动选题：同义词合并后 n >= 5 的高频主题，剔除泛词与主题词本身
    mapper = kw_mapper(df)
    generic = {
        'humans', 'male', 'female', 'middle aged', 'adult', 'aged', 'child',
        'young adult', 'adolescent', 'aged, 80 and over', 'retrospective studies',
        'history, 19th century', 'history, 20th century', 'case reports',
    }
    counts = Counter()
    for val in df['keywords']:
        counts.update(map_terms(keywords(val), mapper))
    tracked = [t for t, n in counts.most_common()
               if n >= 5 and t.lower() not in generic
               and not t.lower().startswith('false localiz')][:10]
    ranges = []
    for term in tracked:
        years = [int(y) for y, val in zip(df['year'], df['keywords'])
                 if pd.notna(y) and term in map_terms(keywords(val), mapper)]
        if years: ranges.append((term, min(years), max(years), len(years)))
    ranges.sort(key=lambda x: (x[1], -x[3]))
    years = list(range(int(df.year.min()), int(df.year.max()) + 1))
    fig, ax = plt.subplots(figsize=(12.3, 6.4)); ax.set_facecolor('white')
    y_pos = np.arange(len(ranges))[::-1]
    for y, (term, start, end, freq) in zip(y_pos, ranges):
        ax.hlines(y, years[0], years[-1], color=BLUE, lw=2, alpha=.65)
        ax.hlines(y, start, end, color=RED, lw=5)
        ax.scatter([start, end], [y, y], color=RED, s=20, zorder=3)
        ax.text(years[0] - 5, y, term, ha='right', va='center', fontsize=9)
        ax.text(years[-1] + 1.5, y, f'n={freq}', va='center', fontsize=8)
    ax.set_yticks([]); ax.set_xlim(years[0] - 48, years[-1] + 12); ax.set_ylim(-1, len(ranges))
    ax.set_xlabel('Year'); ax.set_title('Research-topic time distribution in false localizing sign literature', loc='left', weight='bold', fontsize=14)
    ax.text(0, -0.14, 'Red segments mark the observed first-to-last publication span for each keyword; they are not citation-burst estimates. Topics arranged chronologically by initial appearance year; only topics with n >= 5 records are shown.', transform=ax.transAxes, fontsize=8)
    for side in ['top', 'right', 'left']: ax.spines[side].set_visible(False)
    return save(fig, 'fig11_topic_time_distribution.png', facecolor='white')


def make_analysis_overview(df):
    annual = df.groupby('year').size().sort_index()
    G, counts = network_data(df)
    fig = plt.figure(figsize=(13, 7.2)); gs = gridspec.GridSpec(2, 2, figure=fig, hspace=.48, wspace=.3)
    ax = fig.add_subplot(gs[0, 0]); ax.bar(annual.index, annual.values, color='#457b9d', width=.9); ax.set_title('(A) Annual publication output', loc='left', weight='bold'); ax.set_xlabel('Year'); ax.set_ylabel('Records')
    ax = fig.add_subplot(gs[0, 1]); top = counts.most_common(8); ax.barh([x[0][:28] for x in top][::-1], [x[1] for x in top][::-1], color='#2a9d8f'); ax.set_title('(B) Most frequent keywords', loc='left', weight='bold'); ax.set_xlabel('Frequency')
    ax = fig.add_subplot(gs[1, :]); pos = nx.spring_layout(G, seed=9, k=1.5, iterations=200); nx.draw_networkx_edges(G, pos, alpha=.28, edge_color='#777', ax=ax); nx.draw_networkx_nodes(G, pos, node_size=[80+60*G.nodes[n]['frequency'] for n in G.nodes], node_color='#e76f51', edgecolors='white', ax=ax); nx.draw_networkx_labels(G, pos, font_size=7, ax=ax); ax.margins(.2); ax.set_title('(C) Keyword co-occurrence structure', loc='left', weight='bold'); ax.axis('off')
    fig.suptitle('Bibliometric analysis overview: false localizing signs', fontsize=16, weight='bold')
    return save(fig, 'fig12_bibliometric_analysis_overview.png')


def zh(run):
    run.font.name = 'Microsoft YaHei'; run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def add_heading(doc, text, level):
    p = doc.add_heading(level=level); zh(p.add_run(text)); return p


def add_text(doc, text):
    p = doc.add_paragraph(); run = p.add_run(text); zh(run); return p


def add_image(doc, path, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path), width=Inches(6.15))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = p.add_run(caption); run.italic = True; run.font.size = Pt(9); zh(run)


def build_doc(paths):
    doc = Document(SOURCE_DOCX)
    doc.add_page_break()
    add_heading(doc, '附录 B  基于原始数据的补充文献计量图表', 1)
    add_text(doc, '本附录借鉴参考图片的表格、网络聚类与时间线呈现方式，使用本研究纳入的 159 篇假定位体征文献重新绘制。所有数值均来自本报告的原始题录数据。由于当前数据不包含被引次数及被引参考文献，未开展高被引文献、共被引或引文突现分析，也未以这些名称标注图表。')
    add_heading(doc, 'B.1 代表性近年文献与核心传播载体', 2)
    add_text(doc, '表 B1 按发表年份列出近年的代表性记录，用于展示研究对象的近期临床与影像学关注方向；该表不代表引文影响力排序。表 B2 和表 B3 分别展示高产期刊和作者。')
    add_image(doc, paths[0], 'Table B1. Ten recent representative records on false localizing signs.')
    add_image(doc, paths[1], 'Table B2. Top 10 journals in false localizing sign research.')
    add_image(doc, paths[2], 'Table B3. Top 10 productive authors in false localizing sign research.')
    add_heading(doc, 'B.2 关键词结构与主题时间分布', 2)
    add_text(doc, '关键词网络按同一篇文献中共同出现的 MeSH 词和其他关键词构建；聚类用于直观呈现主题关联。时间图展示关键词在本数据中首次至最后一次出现的发表时间跨度，红色区段不等同于引文突现。')
    add_image(doc, paths[3], 'Fig. B1. Keyword co-occurrence network and cluster visualization.')
    add_image(doc, paths[4], 'Fig. B2. Research-topic time distribution based on observed publication years.')
    add_image(doc, paths[5], 'Fig. B3. Bibliometric analysis overview based on the present dataset.')
    doc.save(DEST)


def main():
    WORK.mkdir(parents=True, exist_ok=True)
    style()
    df = pd.read_csv(DATA)
    paths = [make_representative_records(df), make_journal_table(), make_author_table(df), make_network_and_clusters(df), make_timeline(df), make_analysis_overview(df)]
    build_doc(paths)
    print(DEST)


if __name__ == '__main__':
    main()
