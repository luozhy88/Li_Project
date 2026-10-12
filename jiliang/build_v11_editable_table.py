#!/usr/bin/env python3
"""Create v11 with an editable three-line Word table for Table B2."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path('/Users/apple/Desktop/github.local/Li_Project/jiliang')
SOURCE = ROOT / 'output/report/10_False_Localizing_Sign_文献计量分析报告_Medicine投稿修订版v10.docx'
DEST = ROOT / 'output/report/10_False_Localizing_Sign_文献计量分析报告_Medicine投稿修订版v11.docx'
STATS = ROOT / 'output/logs/b1_own_stats.json'


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        tc_pr.append(margins)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = margins.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in('w:tcW')
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(width_dxa))
    tc_w.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')

    tbl_ind = tbl_pr.first_child_found_in('w:tblInd')
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), '120')
    tbl_ind.set(qn('w:type'), 'dxa')

    layout = tbl_pr.first_child_found_in('w:tblLayout')
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tbl_pr.append(layout)
    layout.set(qn('w:type'), 'fixed')

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)


def set_three_line_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    # Three-line format: table top, header separator, table bottom only.
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = borders.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            borders.append(node)
        node.set(qn('w:val'), 'nil')
    borders.find(qn('w:top')).set(qn('w:val'), 'single')
    borders.find(qn('w:top')).set(qn('w:sz'), '12')
    borders.find(qn('w:bottom')).set(qn('w:val'), 'single')
    borders.find(qn('w:bottom')).set(qn('w:sz'), '12')

    header = table.rows[0]
    for cell in header.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        cell_borders = tc_pr.first_child_found_in('w:tcBorders')
        if cell_borders is None:
            cell_borders = OxmlElement('w:tcBorders')
            tc_pr.append(cell_borders)
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '8')
        bottom.set(qn('w:space'), '0')
        bottom.set(qn('w:color'), '000000')
        cell_borders.append(bottom)


def format_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5):
    cell.text = ''
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run.font.size = Pt(size)
    run.bold = bold


def insert_before(parent, element, reference):
    parent.insert(parent.index(reference), element)


def main():
    stats = json.loads(STATS.read_text(encoding='utf-8'))
    clusters = stats['clustering']['clusters']
    doc = Document(SOURCE)

    caption = next(p for p in doc.paragraphs if p.text.strip().startswith('Table B2.'))
    old_image_p = caption._p.getprevious()
    drawing_tag = qn('w:drawing')
    while old_image_p is not None and not old_image_p.findall('.//' + drawing_tag):
        old_image_p = old_image_p.getprevious()
    if old_image_p is None:
        raise RuntimeError('Table B2 image paragraph not found.')

    parent = caption._p.getparent()
    parent.remove(old_image_p)
    caption.text = 'Table B2. Data-driven keyword clusters corresponding to Fig. B2.'
    for run in caption.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.italic = True

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = [900, 800, 2200, 5460]
    set_table_geometry(table, widths)
    headers = ['Cluster', 'Nodes', 'LLR label', 'Representative keywords (top members)']
    for cell, header in zip(table.rows[0].cells, headers):
        format_cell(cell, header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)

    for cluster in clusters:
        cells = table.add_row().cells
        values = [
            f"#{cluster['id']}",
            str(cluster['size']),
            cluster['label'],
            '; '.join(cluster['members_top']),
        ]
        for i, (cell, value) in enumerate(zip(cells, values)):
            format_cell(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if i < 2 else WD_ALIGN_PARAGRAPH.LEFT, size=9 if i != 3 else 8.5)
    set_table_geometry(table, widths)
    set_three_line_borders(table)

    # Insert caption, table, and note as one paired block.
    table_element = table._tbl
    parent.remove(table_element)
    insert_before(parent, table_element, caption._p)
    parent.remove(caption._p)
    insert_before(parent, caption._p, table_element)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.space_before = Pt(3)
    note.paragraph_format.space_after = Pt(6)
    run = note.add_run(
        f"Note: same 159 records and the same pruned network used for Fig. B2 "
        f"(70 nodes, 206 links); greedy-modularity clustering with LLR labels; "
        f"Q = {stats['clustering']['modularity_Q']:.3f}, weighted mean S = {stats['clustering']['weighted_mean_silhouette_S']:.3f}."
    )
    run.font.name = 'Times New Roman'
    run.font.size = Pt(8.5)
    run.italic = True
    # Move note after the table and before the following content.
    parent.remove(note._p)
    parent.insert(parent.index(table_element) + 1, note._p)

    doc.save(DEST)
    print(DEST)


if __name__ == '__main__':
    main()
