"""Build an explicitly unfinished research manuscript using the managed DOCX runtime."""
import pathlib, json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=pathlib.Path(__file__).resolve().parents[1]
OUT=ROOT/'output/03_SCI英文工作稿';OUT.mkdir(exist_ok=True)
records={r['pmid']:r for r in json.loads((ROOT/'output/02_PubMed初步分析/records.json').read_text())}
refs=['National Library of Medicine. Muscle Spasticity. MeSH Descriptor D009128 (2026). https://meshb.nlm.nih.gov/record/ui?ui=D009128. Accessed 11 October 2026.']
pmids=['25192034','25628750','28455191','36062002','29986902','35245503','35837181','42527851']
for pmid in pmids:
    r=records[pmid]
    names='; '.join(r['authors'][:6])+('; et al.' if len(r['authors'])>6 else '.')
    details=r['year']+(';' + r['volume'] if r['volume'] else '')+('('+r['issue']+')' if r['issue'] else '')+(':'+r['pages'] if r['pages'] else '')
    refs.append(f"{names} {r['title']} {r['journal']}. {details}. https://doi.org/{r['doi']}. PMID: {pmid}.")
refs.extend([
    'Liu J; Chen R; Huang Q; Cheng H; Cai L; Liu F. Acupuncture for post-stroke spastic hemiplegia: A dual-dimensional evidence synthesis combining bibliometric and meta-analysis. Journal of Family Medicine and Primary Care. 2026;15(5):1873–1879. https://doi.org/10.4103/jfmpc.jfmpc_1889_25. PMID: 42483336.',
    'Chen C. CiteSpace II: Detecting and visualizing emerging trends and transient patterns in scientific literature. Journal of the American Society for Information Science and Technology. 2006;57(3):359–377. https://doi.org/10.1002/asi.20317.',
    'Chen C; Ibekwe-SanJuan F; Hou J. The structure and dynamics of cocitation clusters: A multiple-perspective cocitation analysis. Journal of the American Society for Information Science and Technology. 2010;61(7):1386–1409. https://doi.org/10.1002/asi.21309.',
    'Rethlefsen ML; Kirtley S; Waffenschmidt S; et al. PRISMA-S: an extension to the PRISMA Statement for Reporting Literature Searches in Systematic Reviews. Systematic Reviews. 2021;10:39. https://doi.org/10.1186/s13643-020-01542-z.',
    'National Library of Medicine. PubMed user guide. https://pubmed.ncbi.nlm.nih.gov/help/. Accessed 11 October 2026.',
    'Clarivate. Web of Science: Export Records. https://webofscience.help.clarivate.com/en-us/Content/export-records.htm. Accessed 11 October 2026.'
])
assert len(refs)==15
body=(ROOT/'script/manuscript_body.md').read_text()
markdown=body+'\n\n## References\n\n'+'\n\n'.join(f'[{i}] {v}' for i,v in enumerate(refs,1))+'\n'
(OUT/'01_英文论文工作稿_待WoS与正式分析.md').write_text(markdown)
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492)
tokens={'preset':'narrative_proposal','header_pattern':'memo_masthead',
    'named_override':'Academic manuscript: Times New Roman throughout; black headings; compact 18pt title; no decorative rules.',
    'page':'Letter','margins_inches':1,'header_footer_inches':.492,
    'body':{'size_pt':11,'before_pt':0,'after_pt':8,'line_spacing':1.333},
    'headings':{'Heading 1':[16,18,10],'Heading 2':[13,12,6],'Heading 3':[12,8,4]},
    'title':[18,0,10], 'references':{'size_pt':9.5,'after_pt':4,'line_spacing':1.1,'numbering':'real OOXML decimal [%1]'},
    'tables':'none','lists':'only reference numbering','running_header':'WORKING DRAFT | Analysis incomplete'}
def style(name,size,before,after,line=1.333,bold=False):
    s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('000000');s.font.bold=bold
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC')
    rf=s._element.get_or_add_rPr().rFonts
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        if qn('w:'+attr) in rf.attrib:del rf.attrib[qn('w:'+attr)]
    for border in s._element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
    pf=s.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=line
    pf.widow_control=True
    return s
style('Normal',11,0,8).paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
style('Title',18,0,10,1.1,True).paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
style('Subtitle',11,0,8,1.1)
for name,(size,before,after) in tokens['headings'].items():
    hs=style(name,size,before,after,bold=True);hs.paragraph_format.keep_with_next=True;hs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
for name,size in [('Header',9),('Footer',9)]:style(name,size,0,0,1.0)
doc.styles.add_style('Reference Entry',1)
s=style('Reference Entry',9.5,0,4,1.1);s.paragraph_format.left_indent=Inches(.25);s.paragraph_format.first_line_indent=Inches(-.25)
header=sec.header.paragraphs[0];header.text='WORKING DRAFT  |  Acupuncture, spasm and spasticity  |  Analysis incomplete'
header.runs[0].font.color.rgb=RGBColor.from_string('666666')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Page ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for chunk in body.strip().split('\n\n'):
    if chunk.startswith('# '):doc.add_paragraph(chunk[2:],style='Title')
    elif chunk.startswith('## '):doc.add_paragraph(chunk[3:],style='Heading 1')
    elif chunk.startswith('### '):doc.add_paragraph(chunk[4:],style='Heading 2')
    else:doc.add_paragraph(chunk.replace('\n',' '))
doc.add_paragraph('References',style='Heading 1')
numroot=doc.part.numbering_part.element
absid=max([int(n.get(qn('w:abstractNumId'))) for n in numroot.findall(qn('w:abstractNum'))]+[0])+1
numid=max([int(n.get(qn('w:numId'))) for n in numroot.findall(qn('w:num'))]+[0])+1
an=OxmlElement('w:abstractNum');an.set(qn('w:abstractNumId'),str(absid))
level=OxmlElement('w:lvl');level.set(qn('w:ilvl'),'0')
for tag,val in [('start','1'),('numFmt','decimal'),('lvlText','[%1]'),('lvlJc','left')]:
    e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);level.append(e)
pp=OxmlElement('w:pPr');ind=OxmlElement('w:ind');ind.set(qn('w:left'),'360');ind.set(qn('w:hanging'),'360');pp.append(ind)
tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'num');tab.set(qn('w:pos'),'360');tabs.append(tab);pp.append(tabs)
level.append(pp);an.append(level);numroot.append(an)
n=OxmlElement('w:num');n.set(qn('w:numId'),str(numid));a=OxmlElement('w:abstractNumId');a.set(qn('w:val'),str(absid));n.append(a);numroot.append(n)
for ref in refs:
    p=doc.add_paragraph(ref,style='Reference Entry');np=OxmlElement('w:numPr')
    for tag,val in [('ilvl','0'),('numId',str(numid))]:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);np.append(e)
    p._p.get_or_add_pPr().append(np)
doc.core_properties.title='Acupuncture-related therapies for spasm and spasticity — working manuscript'
doc.core_properties.subject='Incomplete bibliometric study; PubMed retrieval audited; WoS and CiteSpace results pending'
doc.core_properties.author=''
dest=OUT/'02_英文论文工作稿_待WoS与正式分析.docx';doc.save(dest)
qa=OUT/'qa';qa.mkdir(exist_ok=True)
(qa/'document_style.json').write_text(json.dumps(tokens,indent=2))
check=Document(dest)
assert check.sections[0].page_width==Inches(8.5)
assert check.sections[0].left_margin==Inches(1)
assert len([p for p in check.paragraphs if p.style.name=='Reference Entry'])==15
assert not check.tables
(qa/'structural_validation.json').write_text(json.dumps({'paragraphs':len(check.paragraphs),'references':15,'tables':0,'page_geometry':'passed','render_status':'pending'},indent=2))
library='# 已核实的背景与方法学文献\n\n这些是写作参考来源，不代表最终纳入清单；英文稿引用顺序与此处一致。前期收集以题录、摘要和可用公开全文为依据，未声称下载所有全文。\n\n'+'\n\n'.join(f'{i}. {v}' for i,v in enumerate(refs,1))
(ROOT/'output/01_检索与研究方案/02_已核实背景与方法学文献.md').write_text(library)
print(dest)
print('Words (body):',len(body.split()),'; references:',len(refs))
