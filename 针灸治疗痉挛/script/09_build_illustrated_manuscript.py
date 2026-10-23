"""Create traceable exploratory figures and insert them next to their discussion."""
import os
os.environ.setdefault('MPLCONFIGDIR','/tmp/mpl_acupuncture_illustrated')
import collections
import copy
import csv
import hashlib
import itertools
import json
import re
import textwrap
import unicodedata
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import numpy as np
import networkx as nx
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/07_正文图文增强版'
FIG=OUT/'正文插图'; DATA=OUT/'图件源数据'
FIG.mkdir(parents=True,exist_ok=True);DATA.mkdir(exist_ok=True)
rs=json.loads((ROOT/'output/02_PubMed初步分析/records.json').read_text())
kw=json.loads((ROOT/'output/06_Medicine审稿修订/核查证据/05_作者关键词与其他词源分离.json').read_text())
candidates=[r for r in rs if r['screen_status']!='provisional_exclude']
assert len(rs)==395 and len(candidates)==193
assert all('eng' in r['languages'] for r in candidates)
T='#167D8D'; B='#28547B'; O='#D58A3A'; GRAY='#687584'; LIGHT='#E9F2F4'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False,'axes.titleweight':'bold','svg.fonttype':'none','pdf.fonttype':42})
manifest=[]

def csvsave(name,rows,fields=None):
    with (DATA/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields or list(rows[0]));w.writeheader();w.writerows(rows)

def save(fig,num,name,scope):
    stem=f'图{num}_{name}'
    for ext in ['png','svg','pdf']:
        fig.savefig(FIG/f'{stem}.{ext}',dpi=300,bbox_inches='tight',facecolor='white')
    plt.close(fig)
    manifest.append({'number':num,'file':stem+'.png','scope':scope,'sha256':hashlib.sha256((FIG/(stem+'.png')).read_bytes()).hexdigest()})

def box(ax,x,y,w,h,label,color=LIGHT,edge=T,fs=11):
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.014,rounding_size=0.018',fc=color,ec=edge,lw=1.3))
    ax.text(x+w/2,y+h/2,label,ha='center',va='center',fontsize=fs,color='#183449',linespacing=1.5)

def arrow(ax,a,b,color=GRAY):
    ax.annotate('',xy=b,xytext=a,arrowprops={'arrowstyle':'-|>','color':color,'lw':1.5,'shrinkA':2,'shrinkB':2})

# 1. Actual retrieved records and explicitly unfinished screening.
fig,ax=plt.subplots(figsize=(10.5,6.1));ax.set(xlim=(0,1),ylim=(0,1));ax.axis('off')
ax.set_title('Retrieval and provisional screening status',loc='left',pad=16)
box(ax,.02,.75,.27,.16,'Baseline PubMed search\nn = 395')
box(ax,.365,.75,.27,.16,'Expanded PubMed search\nn = 604')
box(ax,.71,.75,.27,.16,'Additional identifiers\nn = 209',color='#FFF3E3',edge=O)
arrow(ax,(.295,.83),(.355,.83));arrow(ax,(.64,.83),(.7,.83))
ax.text(.33,.93,'395 retained',ha='center',fontsize=9,color=GRAY)
box(ax,.02,.45,.27,.15,'Provisional exclusions\nn = 202',color='#F0F2F4',edge=GRAY)
box(ax,.365,.45,.27,.15,'Not provisionally excluded\nn = 193')
arrow(ax,(.155,.75),(.155,.61));arrow(ax,(.29,.76),(.49,.61))
box(ax,.02,.18,.27,.16,'Relevance review pending\nn = 182')
box(ax,.365,.18,.27,.16,'Provisionally relevant\nn = 11 (7 + 4)')
arrow(ax,(.46,.45),(.155,.35));arrow(ax,(.50,.45),(.50,.35))
box(ax,.71,.22,.27,.38,'WoS export: pending\nCross-database merge: pending\nFinal eligibility: pending\nFinal included count: unknown',color='#FFF3E3',edge=O,fs=10)
arrow(ax,(.64,.255),(.70,.38))
ax.text(.02,.035,'Provisional workflow, not a completed PRISMA diagram. The 209 additional records are not yet screened.',fontsize=10,color=GRAY)
save(fig,1,'检索与初筛流程','395 original records; 604 expanded identifiers; incomplete screening')

# 2. Study-specific content extraction diagram, no observed counts.
fig,ax=plt.subplots(figsize=(10.5,5.4));ax.axis('off');ax.set(xlim=(0,1),ylim=(0,1))
ax.set_title('Disease-stratified content analysis: planned workflow',loc='left',pad=15)
cols=[(.025,'DISEASE / MODEL',['Neurogenic spasticity','Hemifacial spasm','Other spasm contexts','Human / animal study']),(.36,'INTERVENTION',['Manual / electroacupuncture','Moxibustion / fire needle','Dry needling','Comparator + cointerventions']),(.695,'OUTCOME',['Spasticity / muscle tone','Movement / daily activity','Symptoms / safety / follow-up','Mechanistic measurements'])]
for x,title,labels in cols:
    ax.text(x+.14,.88,title,ha='center',color=B,fontweight='bold')
    for i,label in enumerate(labels):box(ax,x,.70-i*.135,.28,.095,label,fs=9.6)
arrow(ax,(.31,.535),(.35,.535));arrow(ax,(.645,.535),(.685,.535))
box(ax,.17,.07,.66,.11,'Verified within-study links → disease-stratified matrices\nSeparate clinical outcomes from animal mechanism measurements',color='#FFF3E3',edge=O,fs=10)
arrow(ax,(.5,.285),(.5,.19))
save(fig,2,'疾病干预结局分类框架','Conceptual analysis plan; no empirical disease/intervention counts')

# 3. Corrected metadata, including candidate coverage needed for keywords.
fig,axs=plt.subplots(1,3,figsize=(11.4,4.3),gridspec_kw={'width_ratios':[1.05,1,1]})
groups=[(['Author\nkeywords','Affiliations'],[207,344],[187,50],394,'A  Journal records (n = 394)'),(['Author\nkeywords'],[110],[83],193,'B  Candidate coverage (n = 193)')]
for ax,(labels,yes,no,total,title) in zip(axs[:2],groups):
    xx=np.arange(len(labels));ax.bar(xx,yes,color=T,label='Available');ax.bar(xx,no,bottom=yes,color='#DDE3E8',label='Missing')
    for i,(a,b) in enumerate(zip(yes,no)):
        ax.text(i,a/2,str(a),ha='center',va='center',color='white',fontweight='bold')
        ax.text(i,a+b/2,str(b),ha='center',va='center',color=GRAY)
    ax.set(xticks=xx,xticklabels=labels,ylim=(0,total*1.1),ylabel='Records');ax.set_title(title,fontsize=10,loc='left')
axs[0].legend(frameon=False,fontsize=9,loc='upper right')
axs[2].barh(['Publisher author\nkeywords','NLM old-index\nterms only','No KeywordList'],[207,1,186],color=[T,O,'#BEC9D2'])
for i,n in enumerate([207,1,186]):axs[2].text(n+4,i,str(n),va='center',fontsize=10)
axs[2].invert_yaxis();axs[2].set(xlabel='Journal records',xlim=(0,245));axs[2].set_title('C  Keyword source distinction',fontsize=10,loc='left')
fig.suptitle('Metadata coverage before enrichment',x=.02,ha='left',fontweight='bold',fontsize=14)
fig.tight_layout(rect=(0,.04,1,.91));save(fig,3,'书目字段与词源核查','394 journal records and 193 provisional candidates; author keywords corrected')
csvsave('图3_字段覆盖.csv',[{'population':'journal','field':'author_keywords','available':207,'missing':187,'n':394},{'population':'journal','field':'affiliations','available':344,'missing':50,'n':394},{'population':'candidate','field':'author_keywords','available':110,'missing':83,'n':193}])

# 4. Annual distribution using recorded journal publication years, not harmonized final years.
years=list(range(min(int(r['year']) for r in rs),2027))
ac=collections.Counter(int(r['year']) for r in rs);cc=collections.Counter(int(r['year']) for r in candidates)
annual=[{'year':y,'retrieved':ac[y],'candidate':cc[y],'partial_year':y==2026} for y in years]
csvsave('图4_年度候选记录.csv',annual)
fig,axs=plt.subplots(2,1,figsize=(10.5,6.4),gridspec_kw={'height_ratios':[1,1.1]})
for ax,ys,title in [(axs[0],years,'A  Full recorded time range'),(axs[1],list(range(2014,2027)),'B  Recent years (2014–2026)')]:
    ax.bar(ys,[ac[y] for y in ys],color='#C4D8E2',label='Retrieved records (n = 395)')
    ax.plot(ys,[cc[y] for y in ys],'-o',color=T,lw=1.8,ms=3,label='Provisional candidates (n = 193)')
    ax.axvspan(2025.5,2026.5,color=O,alpha=.19);ax.set_ylabel('Records');ax.set_title(title,loc='left',fontsize=11)
axs[0].legend(frameon=False,fontsize=9,loc='upper left');axs[1].set_xticks(range(2014,2027));axs[1].set_xlabel('Publication year recorded in the local PubMed export')
fig.suptitle('Annual distribution of search records — before final eligibility',x=.025,ha='left',fontweight='bold',fontsize=13)
fig.text(.025,.012,'Shading: incomplete 2026. No regression or development-stage classification was performed.',fontsize=10,color=GRAY)
fig.tight_layout(rect=(0,.035,1,.94));save(fig,4,'检索记录年度分布','395 retrieved / 193 provisional candidates; no final publication trend')

# 5. Journal and type metadata, diagnostic only; overlapping type tags retained.
jcounts=collections.Counter(r['journal'] for r in candidates)
topjournals=sorted(jcounts,key=lambda x:(-jcounts[x],x))[:10]
abbrev={'Journal of traditional Chinese medicine = Chung i tsa chih ying wen pan':'J Tradit Chin Med','Evidence-based complementary and alternative medicine : eCAM':'Evid Based Complement Alternat Med','Acupuncture in medicine : journal of the British Medical Acupuncture Society':'Acupunct Med','Journal of alternative and complementary medicine (New York, N.Y.)':'J Altern Complement Med','Archives of physical medicine and rehabilitation':'Arch Phys Med Rehabil','Complementary therapies in medicine':'Complement Ther Med','Journal of acupuncture and meridian studies':'J Acupunct Meridian Stud','Frontiers in neurology':'Front Neurol','Brain and behavior':'Brain Behav'}
types=['Randomized Controlled Trial','Systematic Review','Review','Meta-Analysis','Clinical Trial','Clinical Trial Protocol']
tc={t:sum(t in r['publication_types'] for r in candidates) for t in types}
fig,axs=plt.subplots(1,2,figsize=(11.4,5),gridspec_kw={'width_ratios':[1.15,1]})
axs[0].barh([abbrev.get(j,j) for j in topjournals],[jcounts[j] for j in topjournals],color=T)
for i,j in enumerate(topjournals):axs[0].text(jcounts[j]+.2,i,str(jcounts[j]),va='center',fontsize=9)
axs[0].invert_yaxis();axs[0].set(xlim=(0,max(jcounts.values())+3),xlabel='Candidate records');axs[0].set_title('A  Ten most frequent source journals',loc='left',fontsize=10)
type_labels=['Randomized trial','Systematic review','Review','Meta-analysis','Clinical trial','Trial protocol']
axs[1].barh(type_labels,[tc[t] for t in types],color=B)
for i,t in enumerate(types):axs[1].text(tc[t]+.4,i,str(tc[t]),va='center',fontsize=9)
axs[1].invert_yaxis();axs[1].set(xlim=(0,max(tc.values())+7),xlabel='Records with the PubMed tag');axs[1].set_title('B  Selected publication-type tags',loc='left',fontsize=10)
fig.suptitle('Source and document-type profile of 193 provisional candidates',x=.015,ha='left',fontweight='bold',fontsize=13)
fig.text(.015,.015,'Candidate-set diagnostics only. Type tags overlap; categories do not sum to 193.',fontsize=10,color=GRAY)
fig.tight_layout(rect=(0,.055,1,.93));save(fig,5,'候选文献期刊与类型','193 not provisionally excluded records; journal ranking not final; non-exclusive tags')
csvsave('图5_全部候选期刊.csv',[{'journal':j,'n':jcounts[j],'shown_in_top10':j in topjournals} for j in sorted(jcounts,key=lambda j:(-jcounts[j],j))])
csvsave('图5_类型标签.csv',[{'type':t,'n':tc[t],'denominator':193,'overlap_allowed':True} for t in types])

# 6 / 7. Author-keyword diagnostics, preserving clinically distinct phrases.
def normalize(term):return re.sub(r'\s+',' ',unicodedata.normalize('NFKC',term).strip().lower())
docs={r['pmid']:sorted(set(normalize(t) for t in kw.get(r['pmid'],{}).get('author_keywords',[]) if normalize(t))) for r in candidates}
docs={p:terms for p,terms in docs.items() if terms}
assert len(docs)==110
freq=collections.Counter(t for terms in docs.values() for t in terms)
terms=sorted(freq,key=lambda t:(-freq[t],t))[:20]
edges=collections.Counter(pair for ts in docs.values() for pair in itertools.combinations(sorted(set(ts)&set(terms)),2))
G=nx.Graph();G.add_nodes_from(terms)
G.add_weighted_edges_from((a,b,n) for (a,b),n in sorted(edges.items()) if n>=3)
pos=nx.circular_layout(terms)
fig,ax=plt.subplots(figsize=(10.5,7));ax.set_aspect('equal');ax.axis('off')
nx.draw_networkx_edges(G,pos,ax=ax,width=[.55+.11*G[a][b]['weight'] for a,b in G.edges()],edge_color='#7FA9B3',alpha=.5)
nx.draw_networkx_nodes(G,pos,ax=ax,node_size=[110+18*freq[t] for t in G],node_color=[freq[t] for t in G],cmap='YlGnBu',edgecolors='white',linewidths=1.3)
for t,(x,y) in pos.items():
    label=textwrap.fill(t,21)+f' ({freq[t]})'
    ax.text(x*1.12,y*1.12,label,fontsize=9,ha='left' if x>.12 else 'right' if x<-.12 else 'center',va='center',bbox={'facecolor':'white','alpha':.85,'edgecolor':'none','pad':1})
ax.set(xlim=(-1.95,1.95),ylim=(-1.32,1.32))
ax.set_title('Author-keyword co-occurrence in 110 keyword-bearing candidates',loc='left',fontsize=12,pad=18)
fig.text(.03,.01,'Top 20 terms; edges ≥ 3 co-occurring papers. Circular layout is not a distance or cluster estimate.\nExploratory candidate map, not a CiteSpace result or a final thematic network.',fontsize=10,color=GRAY)
save(fig,6,'候选作者关键词共现','110 keyword-bearing provisional candidates; NetworkX circular layout; top20; edges>=3')
csvsave('图6_关键词全部频次.csv',[{'keyword':t,'paper_count':freq[t],'shown':t in terms} for t in sorted(freq,key=lambda t:(-freq[t],t))])
csvsave('图6_关键词边.csv',[{'term_a':a,'term_b':b,'cooccurring_papers':n,'shown':n>=3} for (a,b),n in sorted(edges.items())])
csvsave('图6_逐篇规范化关键词.csv',[{'pmid':p,'keywords':json.dumps(ts,ensure_ascii=False)} for p,ts in docs.items()])
csvsave('关键词规范化字典.csv',[{'original':t,'normalized':normalize(t)} for t in sorted(set(t for r in candidates for t in kw.get(r['pmid'],{}).get('author_keywords',[])))])
lookup={r['pmid']:r for r in candidates};kwyears=list(range(2014,2027));selected=terms[:15]
byyear=collections.Counter(int(lookup[p]['year']) for p in docs)
matrix=np.array([[sum(t in ts and int(lookup[p]['year'])==y for p,ts in docs.items()) for y in kwyears] for t in selected])
fig,axs=plt.subplots(2,1,figsize=(10.5,6.8),gridspec_kw={'height_ratios':[.6,4]},sharex=True)
axs[0].bar(range(len(kwyears)),[byyear[y] for y in kwyears],color=T);axs[0].set_ylabel('N',rotation=0,labelpad=12)
for i,y in enumerate(kwyears):axs[0].text(i,byyear[y]+.4,str(byyear[y]),ha='center',fontsize=9)
axs[0].set_ylim(0,max(byyear.values())*1.35);axs[0].spines['bottom'].set_visible(False)
im=axs[1].imshow(matrix,cmap='YlGnBu',aspect='auto',vmin=0)
axs[1].set_yticks(range(len(selected)),selected);axs[1].set_xticks(range(len(kwyears)),[str(y)+'*' if y==2026 else str(y) for y in kwyears],rotation=45,ha='right')
for i in range(len(selected)):
    for j in range(len(kwyears)):
        if matrix[i,j]>0:axs[1].text(j,i,str(matrix[i,j]),ha='center',va='center',fontsize=8,color='white' if matrix[i,j]>matrix.max()*.55 else '#253E4F')
fig.suptitle('Annual occurrence of frequent author keywords — candidate subset',x=.02,ha='left',fontsize=13,fontweight='bold')
fig.text(.02,.015,'Cells: raw paper counts, not burst strengths. Top bars: keyword-bearing papers each year.\nOnly 110 / 193 candidates have author keywords; *2026 is incomplete.',fontsize=10,color=GRAY)
fig.tight_layout(rect=(0,.065,1,.94));save(fig,7,'候选关键词年度热图','110 records dated2014–2026; raw frequency; no burst or phase inference')
csvsave('图7_年度关键词计数.csv',[{'keyword':t,'year':y,'n':int(matrix[i,j]),'keyword_bearing_papers_in_year':byyear[y]} for i,t in enumerate(selected) for j,y in enumerate(kwyears)])
csvsave('全部图件候选记录清单.csv',[{'pmid':r['pmid'],'year':r['year'],'status':r['screen_status'],'has_author_keywords':r['pmid'] in docs,'journal':r['journal']} for r in candidates])

# Preserve the reviewed text and its superscript citations; insert figures locally.
doc=Document(ROOT/'output/06_Medicine审稿修订/01_针灸治疗痉挛_Medicine审稿修订稿.docx')
doc.sections[0].header.paragraphs[0].text='针灸相关干预与痉挛｜正文图文阅读版｜候选集探索性分析'
doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].paragraph_format.line_spacing=1.15
doc.styles['Caption'].paragraph_format.space_after=Pt(9)

def find(prefix):return next(p for p in doc.paragraphs if p.text.startswith(prefix))
def after(anchor,text='',style=None):
    p=doc.add_paragraph(text,style);anchor._p.addnext(p._p);return p
def figure(anchor,num,caption):
    anchor.paragraph_format.keep_with_next=True
    anchor.paragraph_format.keep_together=True
    p=after(anchor);p.alignment=1;p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1
    obj=p.add_run().add_picture(str(FIG/manifest[num-1]['file']),width=Inches(6.8))
    obj._inline.docPr.set('descr',caption)
    cap=after(p,caption,'Caption');cap.paragraph_format.keep_together=True
    return cap

find('审稿修订版｜').text='正文图文阅读版｜2026年10月14日｜7幅正文图及1张正文表'
find('本稿按Medicine').text='本版按阅读顺序将图表嵌入方法及结果附近，便于逐段核对。数据图来自本项目保存的检索记录及待复核候选集，方法框架图标为计划。双库筛选及正式分析尚未完成；本版仍为中文工作稿。作者信息待填写。'
p=find('结果：原式检索')
p.text += '对193条待复核候选记录补充期刊及文献类型描述，并对其中110条具有作者关键词的记录绘制探索性共现图和年度热图。'
p=find('Web of Science计划导出');p.text += '图1展示当前检索与初筛状态。'
figure(p,1,'图1 检索与初筛进度。原式395条记录中202条暂标排除、193条尚待继续核查；193条中182条待相关性判断，11条初步相关（7条符合旧脚本基础字段检查、4条待补全）。扩展式新增209条仅保存标识符。本图不是已完成的PRISMA筛选图，最终纳入数尚未确定。')
p=find('疾病—干预与干预—结局矩阵分别构建');p.text += '图2说明内容编码及关联分析的计划。'
figure(p,2,'图2 按疾病分层的内容分析框架。依据本研究问题组织疾病或模型、干预及结局字段，仅连接原文能够确认的对应关系。图中类别为编码方案，不表示已经观察到的数量、效应或机制通路。')
p=find('本轮使用Python解析');p.text += '本版另对未暂标排除的193条记录进行探索性描述，未改变任何资格判断。关键词仅来自NOTNLM字段，按Unicode兼容规范、大小写及空白规范化，同篇同词只计一次，不合并不同临床概念。前20词按论文频次降序选取，同频时按规范词字母序；共现边至少出现在3篇论文中，以NetworkX圆形布局展示，未进行CiteSpace聚类或突现计算。前15词的年度热图采用原始计数，同时展示每年有关键词的文献数。年份取现有PubMed期刊发表年份，尚未完成首次公开发表日期的统一核查。'
p=find('394条期刊记录中，207条');p.text=p.text.replace('见表1','见图3及表1')
cap=figure(p,3,'图3 字段完整性及关键词来源。A：394条期刊记录的作者关键词与机构字段覆盖；B：193条待复核候选记录中110条有作者关键词、83条缺失；C：全部期刊记录中207条有出版方作者关键词、1条仅有NLM旧版索引词、186条无KeywordList。缺少作者关键词的正确计数为187条。')
tp=after(cap,'表1 当前检索记录及字段检查结果','Caption');tp.paragraph_format.keep_with_next=True
source_table=Document(ROOT/'output/06_Medicine审稿修订/05_表1_本项目检索与字段核查.docx')
table_node=copy.deepcopy(source_table.tables[0]._tbl);tp._p.addnext(table_node)
note=doc.add_paragraph('注：均为检索或初筛阶段数据；不同缺项可重叠，193条并非最终纳入数。作者关键词按NOTNLM来源识别。','Caption');table_node.addnext(note._p)
p=find('### 3.3') if any(p.text.startswith('### 3.3') for p in doc.paragraphs) else find('3.3. 正式分析')
p.text='3.7. 正式分析完成情况'
# Insert all new subsections immediately before the old final-analysis status.
def before_target(text,style=None):
    pnew=doc.add_paragraph(text,style);p._p.addprevious(pnew._p);return pnew
before_target('3.3. 候选记录的年度分布','Heading 2')
q=before_target('图4比较全部395条检索记录与193条待复核候选记录的年度分布，既展示完整时间范围，也放大2014年以来的记录。两条序列对应不同筛选状态，均不代表最终领域发文量。2026年为不完整年份，未据此拟合增长曲线或划分发展阶段。')
figure(q,4,'图4 检索记录与待复核候选记录的年度分布。A：现有记录的完整年份范围；B：2014—2026年局部放大。年份沿用本地PubMed记录，尚未统一核对电子与纸刊日期；橙色区域为不完整的2026年。')
before_target('3.4. 候选文献的期刊与类型构成','Heading 2')
q=before_target(f'在193条待复核候选记录中，Medicine有{jcounts["Medicine"]}条，Frontiers in Neurology有{jcounts["Frontiers in neurology"]}条；图5A展示按当前记录数排序的前10种来源期刊。图5B列出部分PubMed文献类型标签。标签可以重叠，且研究方案仍在候选集中，因此该图用于了解待筛选材料的组成，不能作为正式纳入研究的期刊排名或研究设计比例。')
figure(q,5,'图5 待复核候选记录的来源及文献类型。A：按193条记录计数的前10种期刊，同频时按原始期刊名字母序选取；B：选定文献类型标签的记录数，各标签非互斥。正式筛选与类型复核后需重新统计。')
before_target('3.5. 作者关键词的探索性共现','Heading 2')
q=before_target(f'193条候选记录中，110条（{110/193*100:.2f}%）提供作者关键词。图6展示频次最高的20个规范词及至少在3篇文献中共同出现的连接。acupuncture、stroke和spasticity分别出现于{freq["acupuncture"]}、{freq["stroke"]}和{freq["spasticity"]}条有关键词的候选记录。systematic review等词反映出版或方法特征，不能与疾病、干预词作同义合并。图中连线表示同篇共现，不表示治疗效果、因果关系或经验证的主题聚类。')
figure(q,6,f'图6 待复核候选记录的作者关键词共现图。使用110条有关键词的记录；节点为前20词，大小及颜色深浅对应论文频次，括号为频次；连线为至少3篇论文内的共现，共显示{G.number_of_edges()}条边。圆形布局不表示语义距离，节点颜色不表示聚类。本图由本项目脚本生成，不是CiteSpace正式输出。')
before_target('3.6. 作者关键词的年度出现情况','Heading 2')
q=before_target('图7展示前15个作者关键词的逐年记录数，并在上方列出每年有关键词的候选记录数。这110条记录分布于2014—2026年，早期候选文献的关键词覆盖不足，不能把热图左侧视为该领域研究的起点。颜色仅表示原始计数；样本量变化、字段缺失及尚未完成的资格筛选，均会影响年际比较。')
figure(q,7,'图7 待复核候选记录的作者关键词年度热图。按110条有作者关键词的记录计算，每篇论文同词计一次，选取总体频次前15词；上方柱形为各年有关键词的候选记录数。单元格为论文数而非标准化强度或突现值。2026年以星号标记为不完整年份。')
find('最终纳入集、国家与机构').text='最终纳入集、国家与机构及作者合作网络、正式主题聚类、共被引网络和疾病—干预—结局矩阵尚未形成。图4—图7为本地检索或待复核候选集的探索性描述，完成两库筛选后需重新计算；不能由此确定领域领先主体、正式热点或临床转化程度。'
q=find('文献计量描述研究活动');q.text=q.text.replace('现阶段尚无正式网络或内容分布结果','本版关键词图仅描述待复核候选集，尚无正式网络或内容分布结果')
q=find('当前直接限制是筛选');q.text += '关键词探索图另受83条候选记录缺失作者关键词的影响，110条有词记录不能未经比较就代表全部193条候选，更不能代表最终纳入集。'
doc.core_properties.subject='正文嵌图阅读版；7幅图与1张表；本地数据探索性分析；最终纳入尚未完成'
dest=OUT/'01_针灸治疗痉挛_正文图文增强版.docx';doc.save(dest)
check=Document(dest)
assert len(check.inline_shapes)==7 and len(check.tables)==1
ref_index=next(i for i,x in enumerate(check.paragraphs) if x.text=='参考文献')
assert all(i<ref_index for i,x in enumerate(check.paragraphs) if x._p.xpath('.//w:drawing'))
captions=[x.text for x in check.paragraphs if re.match(r'^图[1-7] ',x.text)]
assert [int(t[1]) for t in captions]==list(range(1,8))
assert sum(1 for x in check.paragraphs if re.match(r'^\d+\. ',x.text) and x.style.name=='Normal')==16
(OUT/'图件清单与核查.json').write_text(json.dumps({'figures':manifest,'candidate_records':193,'keyword_records':110,'keyword_edges_displayed':G.number_of_edges(),'nodes_displayed':20,'heatmap_rows':15,'all_figures_before_references':True,'embedded_tables':1,'software':{'matplotlib':matplotlib.__version__,'networkx':nx.__version__},'final_eligibility_completed':False},ensure_ascii=False,indent=2))
text_lines=[]
for par in check.paragraphs:
    if par._p.xpath('.//w:drawing'):continue
    if par.text:text_lines.append(('## ' if par.style.name=='Heading 1' else '### ' if par.style.name=='Heading 2' else '')+par.text)
(OUT/'02_正文图文版文本.md').write_text('\n\n'.join(text_lines)+'\n')
(OUT/'03_图文调整说明.md').write_text('''# 正文图文调整说明

本版按用户要求制作阅读用Word：7幅图、1张表均放在对应正文附近，未集中置于文末。保留06版已核验的16项参考文献及上标引用；图件附独立PNG、SVG、PDF和可核查数据，正式提交期刊时仍可使用单独图件。

图1为实际检索与初筛进度；图2为疾病—干预—结局编码方案；图3为修正后的书目字段覆盖；图4为检索及候选记录年度分布；图5为候选期刊和文献类型构成；图6为候选作者关键词共现；图7为候选关键词年度热图。每幅图均有正文引述和紧随其后的中文图题，图内使用英文标识以便后续英文定稿。

借鉴参考文章的流程图、分类框架、年度曲线、来源分布和关键词时间图组织方式。所有统计由本项目数据生成，未复制参考文章的图、数字或分类结果。图2是本研究的问题分类示意，不是统计结果。

193条候选仍未完成资格复核，其中只有110条有作者关键词。图4—图7用于探索性检查和阅读讨论，不代表已完成的领域计量结果。未将国家地图、机构合作、共被引、突现强度和疾病—干预桑基图用假设数值补齐：这些图所需的最终纳入、名称规范化、完整引文或内容编码尚未完成。

关键词仅作Unicode兼容、大小写和空白规范化；不合并spasticity、muscle spasticity及post-stroke spasticity等语义相关但并不相同的词。每篇同词计一次；前20词按频次及字母序选取；图6边阈值为共现至少3篇，使用圆形布局，没有声称CiteSpace聚类或计算热点突现。图7为前15词的逐年原始计数，并显示每年词源可用的分母；2014年只是当前关键词可得子集的起始年。

图5期刊按当前候选记录统计，同频按原始期刊名字母序；只列部分PubMed类型标签，标签可重叠，未把相加结果当成总篇数。图4沿用本地PubMed记录年份，未将其伪装成已核查统一后的最终发表年份。

原稿、原始数据及既有筛选状态均保留。生成脚本：script/09_build_illustrated_manuscript.py。图件源数据可按PMID追溯；没有执行新的资格筛选。
''')
print(dest)
print(json.dumps({'figures':7,'candidates':193,'with_author_keywords':110,'top_terms':[(t,freq[t]) for t in terms],'edges':G.number_of_edges()},ensure_ascii=False))
