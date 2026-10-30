from pathlib import Path
import json,csv,re,sys
from collections import Counter
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/11_双库方案修订'
rows=json.loads((ROOT/'output/09_审稿后重分析/screened_records.json').read_text())
c=Counter(x['exclusion_reason'] for x in rows if x['decision']=='exclude')
with (OUT/'03_旧PubMed排除计数_非双库结果.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['旧筛选互斥主原因','篇数','适用说明'])
 for k,v in c.items():w.writerow([k,v,'旧604条PubMed扩展队列；未按本轮严格缺项标准重筛；不可作为双库结果'])
ptags=Counter(t for x in rows for t in set(x.get('publication_types',[])))
(OUT/'02_修改说明与数据缺口.md').write_text('''# 本轮修改与接续输入

已重写双库工作稿，保留原稿供追溯；正文含10幅图的执行位置、5张Top10表和1张聚类表。所有未完成分析均明确标为待核定，未将单库排名移植为双库结果。

## 与旧稿的实质差异

- 从神经性痉挛主分析调整为用户指定的广义spasm/spasticity范围，需重新核查不同痉挛疾病。
- 原式作为主检索，扩展式作为单独敏感性检索；旧604条是原式395条及扩展新增209条的合并，不能全部称作原式命中。
- 英文原始研究和综述均可纳入；先去重、合并补全，仍缺年份、作者、机构或关键词者按本次要求排除。
- 原稿146篇中存在无关键词或无机构记录，因此不能沿用为本次最终集合。
- 修正网络距离、连接强度与中心性的混用；拟合明确完整年度和模型；突现明确参数与适用范围。
- CiteSpace多版本、LLR/LSI、Excel及桑基图工具均须实际执行后填写，未虚写软件使用记录。

## 旧数据核对（仅供追溯）

固定604条PubMed扩展检索记录，既往筛选：415条排除、16条信息不足、173篇广义相关，其中146篇神经性痉挛。上述集合尚未按本轮书目完整性标准重筛。

旧互斥主原因：病例报告57、Letter 8、Comment 1、主题不符48、干预不符29、人群不符8、研究范围不符4、非英文210、Protocol 40、撤稿文献5、勘误3、撤稿通知1、书籍/报告1，共415条。主题/干预/人群/范围不符为89条，但不宜全部改名为“与痉挛无关”。Editorial未作为旧互斥主原因单列，不能由此认定双库社论删除数为0。16条材料不足单列，不包含在415条中。多标签统计不能与互斥主原因混加。

## 接续所需材料

1. WoS Core Collection检索全部批次Plain Text，导出内容“Full Record and Cited References”，放入input/wos/。附实际检索式、检索时间、命中数及引文索引。
2. 若主检索最终采用用户原式，须分别锁定两库实际执行式；敏感性检索单列。保存2026年文献具体发表日期。
3. 合并、去重、补全后重新筛选；逐条记录病例报告、Editorial、Letter及无关标签和互斥原因，核验流程数量守恒。
4. CiteSpace实际安装版本及正式工程/运行结果；逐图保存时间切片、节点筛选、剪枝、阈值与标签字段。未发现可核实的正式运行结果。
5. 期刊引文连接强度、共被引分析须先评估参考文献覆盖；桑基图另记录所用工具及计数方式。

本轮已完成文稿结构与方法修订；双库检索、合并重筛及正式图谱仍待上述数据。旧图表与旧正文未覆盖，避免破坏既往工作。
''',encoding='utf-8')
# Preserve academic styles and page geometry from the retained manuscript.
doc=Document(ROOT/'output/10_审稿后修订稿/05_中文论文_审稿后修订稿.docx')
body=doc._element.body
for child in list(body):
 if child.tag!=qn('w:sectPr'):body.remove(child)
for sec in doc.sections:sec.header.paragraphs[0].text='针灸治疗痉挛知识图谱分析｜双库待补数据工作稿'
sys.path.insert(0,'/Users/apple/.codex/plugins/cache/openai-primary-runtime/documents/26.805.11740/skills/documents/scripts')
from table_geometry import apply_table_geometry,audit_docx_tables
lines=(OUT/'01_针灸治疗痉挛知识图谱分析_双库待补数据稿.md').read_text().splitlines()
i=0
while i<len(lines):
 s=lines[i].strip();i+=1
 if not s:continue
 if s.startswith('|'):
  data=[s]
  while i<len(lines) and lines[i].startswith('|'):data.append(lines[i]);i+=1
  data=[r for r in data if not re.fullmatch(r'[| :\-]+',r)]
  vals=[[x.strip() for x in r.strip('|').split('|')] for r in data]
  t=doc.add_table(rows=0,cols=len(vals[0]));t.autofit=False
  for n,row in enumerate(vals):
   cells=t.add_row().cells
   for cell,value in zip(cells,row):
    cell.text=value
    for p in cell.paragraphs:
     p.paragraph_format.space_after=Pt(3)
     for run in p.runs:run.font.size=Pt(10);run.bold=n==0
   if n==0:t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  widths=[9360//len(vals[0])]*len(vals[0]);widths[-1]+=9360-sum(widths)
  apply_table_geometry(t,widths,table_width_dxa=9360,indent_dxa=120)
  continue
 if s.startswith('# '):doc.add_paragraph(s[2:],'Title')
 elif s.startswith('## '):doc.add_paragraph(s[3:],'Heading 1')
 elif s.startswith('### '):doc.add_paragraph(s[4:],'Heading 2')
 else:
  p=doc.add_paragraph(s)
  if re.match(r'表\d',s) or s.endswith('原式为：'):p.paragraph_format.keep_with_next=True
file=OUT/'01_针灸治疗痉挛知识图谱分析_双库待补数据稿.docx'
doc.save(file)
print(file)
print(audit_docx_tables(file))
