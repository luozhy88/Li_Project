"""Build the complete Chinese article; narrative_proposal with academic typography override."""
from pathlib import Path
import json,sys,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
sys.path.insert(0,'/Users/apple/.codex/plugins/cache/openai-primary-runtime/documents/26.805.11740/skills/documents/scripts')
from table_geometry import apply_table_geometry,audit_docx_tables
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'output/07_中文论文修订稿';data=json.loads((OUT/'manuscript_components.json').read_text())
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492)
for name,size,before,after in [('Normal',11.5,0,8),('Title',20,0,16),('Heading 1',16,18,10),('Heading 2',13,12,6),('Heading 3',12,10,5),('Caption',10,4,8),('Header',9,0,0),('Footer',9,0,0)]:
 s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC' if name in ['Title','Heading 1','Heading 2','Heading 3'] else 'Songti SC')
 p=s.paragraph_format;p.space_before=Pt(before);p.space_after=Pt(after);p.line_spacing=1.333 if name=='Normal' else 1.15;p.widow_control=True
 if name.startswith('Heading'):p.keep_with_next=True
sec.header.paragraphs[0].text='针灸与干针干预痉挛相关疾病｜PubMed文献计量分析'
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run('第 ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f);p.add_run(' 页')
for chunk in data['body'].strip().split('\n\n'):
 if chunk.startswith('# '):doc.add_paragraph(chunk[2:],style='Title')
 elif chunk.startswith('## '):doc.add_paragraph(chunk[3:],style='Heading 1')
 elif chunk.startswith('### '):doc.add_paragraph(chunk[4:],style='Heading 2')
 else:
  p=doc.add_paragraph(chunk.replace('\n',' '))
  if re.match(r'^\[\d+\]',chunk):
   p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_together=True
   for r in p.runs:r.font.size=Pt(10)
widths=[[1400,5160,1200,1600],[7000,1000,1360],[4300,1400,1700,1960],[1500,6500,1360],[4800,1000,1700,1860],[3600,1000,1000,1200,2560]]
doc.add_page_break();doc.add_paragraph('统计表',style='Heading 1')
for i,(title,headers,rows) in enumerate(data['tables']):
 p=doc.add_paragraph(title,style='Caption');p.paragraph_format.keep_with_next=True
 table=doc.add_table(rows=1,cols=len(headers));table.autofit=False
 for c,v in zip(table.rows[0].cells,headers):c.text=str(v)
 for row in rows:
  for c,v in zip(table.add_row().cells,row):c.text=str(v)
 w=widths[i]
 if len(w)!=len(headers):w=[9360//len(headers)]*len(headers);w[-1]+=9360-sum(w)
 apply_table_geometry(table,w,table_width_dxa=9360,indent_dxa=120,cell_margins_dxa={'top':80,'bottom':80,'start':120,'end':120})
 for j,row in enumerate(table.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if j==0:pr.append(OxmlElement('w:tblHeader'))
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
    for r in p.runs:r.font.size=Pt(10);r.bold=j==0
   borders=OxmlElement('w:tcBorders')
   for edge in (['top','bottom'] if j==0 else ['bottom'] if j==len(table.rows)-1 else []):
    b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'6');b.set(qn('w:color'),'444444');borders.append(b)
   c._tc.get_or_add_tcPr().append(borders)
 doc.add_paragraph()
for name,caption in data['figures']:
 doc.add_page_break();p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
 shape=p.add_run().add_picture(str(ROOT/'output/06_完整分析/figures'/f'{name}.png'),width=Inches(6.5));shape._inline.docPr.set('descr',caption)
 doc.add_paragraph(caption,style='Caption')
doc.core_properties.title=data['title'];doc.core_properties.subject='PubMed文献计量分析；169篇纳入文献';doc.core_properties.author=''
p=OUT/'01_中文论文_完整分析修订稿.docx';doc.save(p)
assert len(doc.tables)==6 and len(doc.inline_shapes)==10
assert len([p for p in doc.paragraphs if re.match(r'^\[\d+\]',p.text)])==15
audit_docx_tables(p)
print(p)
