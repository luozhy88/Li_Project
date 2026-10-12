"""Create a Chinese manuscript without modifying the English source artifacts."""
import json, pathlib, re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=pathlib.Path(__file__).resolve().parents[1]
OUT=ROOT/'output/04_中文文稿';OUT.mkdir(exist_ok=True)
QA=OUT/'排版检查';QA.mkdir(exist_ok=True)
src=ROOT/'output/03_SCI英文工作稿/02_英文论文工作稿_待WoS与正式分析.docx'
doc=Document(src)
oldrefs=[p for p in doc.paragraphs if p.style.name=='Reference Entry']
original=[p.text for p in oldrefs]
numid=oldrefs[0]._p.find('w:pPr/w:numPr/w:numId',namespaces=oldrefs[0]._p.nsmap).get(qn('w:val'))
for e in list(doc._element.body):
    if e.tag!=qn('w:sectPr'):doc._element.body.remove(e)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Reference Entry','Header','Footer']:
    s=doc.styles[name]
    rf=s._element.get_or_add_rPr().rFonts
    rf.set(qn('w:eastAsia'),'PingFang SC' if name in ['Title','Heading 1','Heading 2','Heading 3'] else 'Songti SC')
    s.paragraph_format.widow_control=True
doc.styles['Normal'].font.size=Pt(11.5)
doc.styles['Title'].font.size=Pt(20)
doc.styles['Reference Entry'].font.size=Pt(10)
doc.styles['Reference Entry'].paragraph_format.line_spacing=1.2
doc.styles['Reference Entry'].paragraph_format.keep_together=True
doc.styles['Caption'].font.name='Times New Roman';doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption']._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC')
doc.styles['Caption'].paragraph_format.space_before=Pt(4);doc.styles['Caption'].paragraph_format.space_after=Pt(8)
doc.styles['Caption'].paragraph_format.line_spacing=1.2
doc.styles['Caption'].paragraph_format.keep_together=True
sec=doc.sections[0]
sec.header.paragraphs[0].text='针灸治疗痉挛知识图谱分析｜中文工作稿｜正式分析待完成'
footer=sec.footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
body=(ROOT/'script/manuscript_zh.md').read_text().replace('科学引文数据库核心合集','Web of Science核心合集').replace('科学引文数据库入藏号','Web of Science入藏号').replace('以广义指数选择节点','以g指数选择节点')
for chunk in body.strip().split('\n\n'):
    if chunk.startswith('# '):doc.add_paragraph(chunk[2:],style='Title')
    elif chunk.startswith('## '):
        p=doc.add_paragraph(chunk[3:],style='Heading 1')
        if chunk.startswith('## 附录'):p.paragraph_format.page_break_before=True
    elif chunk.startswith('### '):doc.add_paragraph(chunk[4:],style='Heading 2')
    elif chunk.startswith('!['):
        match=re.match(r'!\[(.*?)\]\((.*?)\)',chunk)
        path=(OUT/match.group(2)).resolve()
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(0)
        inline=p.add_run().add_picture(str(path),width=Inches(6.5))
        inline._inline.docPr.set('descr',match.group(1)+'，图内英文标识，中文说明见图下注释。')
    elif chunk.startswith('附图A'):doc.add_paragraph(chunk,style='Caption')
    else:doc.add_paragraph(chunk.replace('\n',' '))
refs=[
 '美国国家医学图书馆．肌肉痉挛状态：2026年医学主题词描述记录D009128．https://meshb.nlm.nih.gov/record/ui?ui=D009128．访问日期：2026年10月11日。',
 'Park SW，Yi SH，Lee JA，等．针灸治疗卒中后痉挛：随机对照试验的荟萃分析．替代与补充医学杂志，2014，20（9）：672–682．https://doi.org/10.1089/acm.2014.0097。',
 'Lim SM，Yoo J，Lee E，等．针灸治疗卒中后痉挛：随机对照试验的系统评价与荟萃分析．循证补充与替代医学，2015：870398．https://doi.org/10.1155/2015/870398。',
 'Cai Y，Zhang CS，Liu S，等．电针治疗卒中后痉挛：系统评价与荟萃分析．物理医学与康复档案，2017，98（12）：2578–2589.e4．https://doi.org/10.1016/j.apmr.2017.03.023。',
 'Xue C，Jiang C，Zhu Y，等．针灸治疗卒中后痉挛的有效性与安全性：系统评价与荟萃分析．神经病学前沿，2022，13：942597．https://doi.org/10.3389/fneur.2022.942597。',
 'Sánchez-Mila Z，Salom-Moreno J，Fernández-de-Las-Peñas C．干针对卒中后痉挛、运动功能和稳定性界限的影响：随机临床试验．医学针灸，2018，36（6）：358–366．https://doi.org/10.1136/acupmed-2017-011568。',
 'Mu JD，Ma LX，Zhang Z，等．针灸通过钾氯共转运体2介导的脊髓γ-氨基丁酸A型受体激活，减轻缺血性卒中后痉挛性肌张力增高大鼠的脊髓反射亢进与运动功能障碍．实验神经病学，2022，354：114027．https://doi.org/10.1016/j.expneurol.2022.114027。',
 'Chen B，Ma Y，Zhong C，等．针灸抑制结肠镜检查期间结肠痉挛的效果与安全性：随机对照试验．胃肠肿瘤学杂志，2022，13（3）：1169–1177．https://doi.org/10.21037/jgo-22-407。',
 'Li S，Li W，Huang J，等．远端电针通过调节脑干与运动皮层的适应不良性可塑性缓解面肌痉挛．运动障碍，2026．https://doi.org/10.1002/mds.70455。',
 'Liu J，Chen R，Huang Q，等．针灸治疗卒中后痉挛性偏瘫：结合文献计量与荟萃分析的双维度证据综合．家庭医学与初级保健杂志，2026，15（5）：1873–1879．https://doi.org/10.4103/jfmpc.jfmpc_1889_25。',
 'Chen C．科学文献中新兴趋势与短暂模式的检测和可视化：CiteSpace第二版．美国信息科学与技术学会杂志，2006，57（3）：359–377．https://doi.org/10.1002/asi.20317。',
 'Chen C，Ibekwe-SanJuan F，Hou J．共被引聚类的结构与动态：多视角共被引分析．美国信息科学与技术学会杂志，2010，61（7）：1386–1409．https://doi.org/10.1002/asi.21309。',
 'Rethlefsen ML，Kirtley S，Waffenschmidt S，等．系统评价文献检索报告规范：系统评价与荟萃分析优先报告条目的检索扩展．系统评价，2021，10：39．https://doi.org/10.1186/s13643-020-01542-z。',
 '美国国家医学图书馆．PubMed用户指南．https://pubmed.ncbi.nlm.nih.gov/help/．访问日期：2026年10月11日。',
 '科睿唯安．Web of Science记录导出说明．https://webofscience.help.clarivate.com/en-us/Content/export-records.htm．访问日期：2026年10月11日。'
]
p=doc.add_paragraph('参考文献',style='Heading 1');p.paragraph_format.page_break_before=True
note='以下文献题名和期刊名采用中文译名，供阅读使用；作者姓名、数据库及软件专名、数字对象标识符和网址保留原文。中文译名不替代原始题录，核查与投稿时请参照同目录中的《原始参考文献核查表》。'
doc.add_paragraph(note)
for ref in refs:
    p=doc.add_paragraph(ref,style='Reference Entry');np=OxmlElement('w:numPr')
    for tag,val in [('ilvl','0'),('numId',str(numid))]:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);np.append(e)
    p._p.get_or_add_pPr().append(np)
doc.core_properties.title='针灸治疗痉挛的研究格局与主题演变：文献计量与知识图谱分析'
doc.core_properties.subject='中文工作稿；正式双数据库知识图谱分析尚未完成'
dest=OUT/'01_针灸治疗痉挛知识图谱分析_中文工作稿.docx';doc.save(dest)
(OUT/'02_针灸治疗痉挛知识图谱分析_中文工作稿.md').write_text(body+'\n\n## 参考文献\n\n'+note+'\n\n'+'\n\n'.join(f'[{i}] {v}' for i,v in enumerate(refs,1))+'\n')
(OUT/'03_原始参考文献核查表.md').write_text('# 原始参考文献核查表\n\n本文件仅供追溯原始外文题录，不属于中文正文。编号与中文稿一致。\n\n'+'\n\n'.join(f'[{i}] {v}' for i,v in enumerate(original,1))+'\n')
tokens=json.loads((ROOT/'output/03_SCI英文工作稿/qa/document_style.json').read_text())
tokens.update({'named_override':'中文学术稿：正文宋体11.5磅，中文标题苹方，标题20磅；西文专名新罗马体；中文页眉页码；参考文献10磅、1.2倍行距。','images':2})
(QA/'样式记录.json').write_text(json.dumps(tokens,ensure_ascii=False,indent=2))
check=Document(dest)
assert len(check.inline_shapes)==2
assert len([p for p in check.paragraphs if p.style.name=='Reference Entry'])==15
assert 'WORKING DRAFT' not in '\n'.join(p.text for p in check.paragraphs)
print(dest)
print('正文汉字数：',len(re.findall(r'[\u4e00-\u9fff]',body)),'；插图：2；参考文献：15')
