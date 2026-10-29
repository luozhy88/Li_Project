"""Render manuscript sources and keep figures with complete legends."""
from pathlib import Path
import subprocess
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

p=Path(__file__).resolve().parents[1]
pandoc='/home/zhiyu/miniconda3/envs/test/bin/pandoc'
names=['01_Manuscript_EN','02_Manuscript_CN','03_Cover_letter','03_Cover_letter_CN',
       '04_CCHTS_submission_guide','06_Supplementary_Figures']
for name in names:
    subprocess.run([pandoc,name+'.md','--from','markdown+superscript','-o',name+'.docx','--standalone'],cwd=p,check=True)
    doc=Document(p/(name+'.docx'))
    for style_name in ['Normal','Body Text','First Paragraph','Title','Heading 1','Heading 2','Heading 3']:
        if style_name in doc.styles:
            style=doc.styles[style_name]
            style.font.name='Times New Roman'
            style.font.size=Pt(10)
            style.font.color.rgb=RGBColor(0,0,0)
            if style_name.startswith('Heading') or style_name=='Title':
                style.font.bold=True
    for section in doc.sections:
        section.top_margin=Inches(.8); section.bottom_margin=Inches(.8)
        para=section.footer.paragraphs[0]; para.alignment=1
        field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); para._p.append(field)
    image_count=0
    for para in doc.paragraphs:
        if name=='01_Manuscript_EN' and para.text in ['Abstract','1. Introduction']:
            para.paragraph_format.page_break_before=True
        if para._p.xpath('.//w:drawing'):
            image_count+=1
            para.paragraph_format.keep_with_next=True
            if name=='06_Supplementary_Figures' and image_count>1:
                para.paragraph_format.page_break_before=True
        if para.text.startswith(('Figure ','Supplementary Figure ','图1','图2','图3','图4','补充图S')):
            para.paragraph_format.keep_together=True
    doc.save(p/(name+'.docx'))
print('Rendered and formatted',len(names),'documents')
