"""Final package consistency and PDF layout checks."""
from pathlib import Path
import json, re, zipfile
import fitz
from docx import Document
from openpyxl import load_workbook

p=Path(__file__).resolve().parents[1]
out=p/'11_final_review'
en=(p/'01_Manuscript_EN.md').read_text()
cn=(p/'02_Manuscript_CN.md').read_text()
checks={}
for f in ['01_Manuscript_EN.md','02_Manuscript_CN.md','03_Cover_letter.md','03_Cover_letter_CN.md']:
    assert (p/f).read_bytes()==(p/'10_editable_sources'/f).read_bytes(), f
checks['source_copies']='PASS'
title=en.splitlines()[0][2:]
abstract=en.split('# Abstract')[1].split('**Keywords')[0]
checks['title_characters']=len(title)
checks['abstract_words_including_headings']=len(abstract.split())
assert len(title)<=120 and len(abstract.split())<=250
refs=en.split('# References')[1]
assert len(re.findall(r'^\[\d+\]',refs,re.M))==26
assert refs.strip()==cn.split('# 参考文献')[1].split('# 补充图')[0].strip()
numbers=[]
for m in re.finditer(r'\[([\d, -]+)\]',en.split('# References')[0]):
    for part in m[1].split(','):
        ns=list(map(int,part.strip().split('-')))
        for n in range(ns[0],ns[-1]+1):
            if n not in numbers: numbers.append(n)
assert numbers==list(range(1,27)),numbers
checks['references_identical_and_cited_in_order']=26
for md in p.glob('*.md'):
    for path in re.findall(r'!\[[^]]*\]\(([^)]+)\)',md.read_text()):
        assert (p/path).is_file(),path
for name,count in [('01_Manuscript_EN',4),('02_Manuscript_CN',7),('06_Supplementary_Figures',3)]:
    doc=Document(p/(name+'.docx'))
    assert len(doc.inline_shapes)==count
    with zipfile.ZipFile(p/(name+'.docx')) as z:
        assert z.testzip() is None
checks['docx_embedded_figures']='4 EN / 7 CN / 3 supplementary; PASS'
wb=load_workbook(p/'05_Supplementary_Tables.xlsx',read_only=True,data_only=True)
assert len(wb.sheetnames)==40
rows=list(wb['T39_replicated_candidates'].values)
assert len(rows)==84
assert all(row[5]>0 and row[7]>0 for row in rows[1:])
checks['workbook']='39 data sheets + README; 83 replicated candidates; no zero adjusted P'
for name in ['01_Manuscript_EN','06_Supplementary_Figures']:
    pdf=fitz.open(out/(name+'.pdf'))
    summary=[]
    contact=fitz.open()
    cell_w,cell_h=330,490
    overview=contact.new_page(width=cell_w*3,height=cell_h*((len(pdf)+2)//3))
    for i,page in enumerate(pdf):
        text=page.get_text()
        summary.append({'page':i+1,'characters':len(text),'images':len(page.get_image_info()),'first_text':text[:110]})
        x=(i%3)*cell_w; y=(i//3)*cell_h
        overview.insert_text((x+5,y+12),str(i+1),fontsize=10)
        overview.show_pdf_page(fitz.Rect(x+4,y+20,x+cell_w-4,y+cell_h-4),pdf,i)
    overview.get_pixmap().save(out/(name+'_contact_sheet.png'))
    checks[name+'_pages']=summary
(out/'package_validation.json').write_text(json.dumps(checks,indent=2,ensure_ascii=False))
print(json.dumps(checks,indent=2,ensure_ascii=False))
