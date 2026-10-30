# -*- coding: utf-8 -*-
"""生成作者专业领域、机构部门与文章贡献的汇总 Word（中英对照）。"""
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

ROWS = [
    ("Tao Li\n（共同第一作者 Co-first author）",
     "针灸推拿\nAcupuncture and Moxibustion",
     "杭州师范大学附属医院 针灸推拿科（浙江杭州）\nDepartment of Acupuncture and Moxibustion, Affiliated Hospital of "
     "Hangzhou Normal University, Hangzhou, Zhejiang",
     "数据收集；数据整理；统计分析\nData collection; data curation; formal analysis"),
    ("Jiacai Li\n（共同第一作者 Co-first author）",
     "医学实践教学\nPractical medical education",
     "上海健康医学院 实践教学中心（上海）\nPractical Teaching Center, Shanghai University of Medicine & Health "
     "Sciences, Shanghai",
     "指导；稿件审阅与修改\nSupervision; writing \u2013 review and editing"),
    ("Yajie Ding\n（共同第一作者 Co-first author）",
     "肿瘤学\nOncology",
     "上海中医药大学附属龙华医院 肿瘤科（上海）\nDepartment of Oncology, Longhua Hospital Affiliated to Shanghai "
     "University of Traditional Chinese Medicine, Shanghai",
     "数据整理；稿件审阅与修改\nData curation; writing \u2013 review and editing\n邮箱：dingyajie2018@163.com"),
    ("Zhiyu Luo",
     "生物技术\nBiotechnology",
     "江苏三黍生物科技有限公司（江苏南通）\nJiangsu Sanshu Biotechnology Co., Ltd., Nantong, Jiangsu",
     "统计分析\nFormal analysis\n利益冲突声明：该公司员工（employee of the company）"),
    ("Xing Zhang\n（通讯作者 Corresponding author）",
     "呼吸系统疾病\nRespiratory disease",
     "上海中医药大学附属曙光医院 呼吸科（上海 201203）\nDepartment of Respiratory Disease, Shuguang Hospital "
     "Affiliated to Shanghai University of Traditional Chinese Medicine, Shanghai 201203",
     "数据整理；稿件审阅与修改\nData curation; writing \u2013 review and editing\n"
     "邮箱：zhangxing2012@shutcm.edu.cn；电话：021-20256375"),
]

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

doc.add_heading("作者信息汇总：专业领域、机构部门与文章贡献", level=1)
doc.add_paragraph("稿件：消瘿汤核心药组在甲状腺癌中的网络药理学及独立转录组验证 / Network pharmacology and "
                  "independent transcriptome validation of core herbs selected from Xiao Ying Tang in thyroid cancer")
doc.add_paragraph("说明：专业领域按手稿所列所在科室/单位性质归纳；贡献与联系方式均取自手稿原文。")

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
for i, h in enumerate(["作者 Author", "专业领域 Field", "机构与部门 Affiliation & Department",
                       "文章贡献 Contributions"]):
    table.rows[0].cells[i].text = h
for row in ROWS:
    cells = table.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = v

doc.add_paragraph("注：Tao Li、Jiacai Li、Yajie Ding 为共同第一作者（contributed equally, share first "
                  "authorship）；全体作者均已审阅并批准最终稿。")

out = "output/20261027_27_before_final_review/09_作者信息与贡献.docx"
doc.save(out)
print("saved:", out)
