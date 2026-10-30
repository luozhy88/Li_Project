# -*- coding: utf-8 -*-
"""生成投稿用结构化「材料与方法 / 结果 / 结论」Word 参考稿（每部分 ≤500 字）。"""
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

METHODS = (
    "本研究采用网络药理学结合独立转录组验证的策略。消瘿汤核心药组（XYT，含莪术、白芍、当归、柴胡、"
    "夏枯草、浙贝母等14味中药）的成分取自TCMSP数据库，映射至PubChem标识符；成分关联基因整合自"
    "BindingDB、BATMAN-TCM（预测评分>0.80）及IUPHAR/BPS Guide to PHARMACOLOGY，与GeneCards的12,987个"
    "甲状腺肿瘤相关基因取交集，分别构建XYT、浙贝母（ZBM）和夏枯草（XKC）候选基因集，并用"
    "clusterProfiler进行GO和KEGG富集分析。表达验证基于两个独立的配对甲状腺乳头状癌队列（GSE60542，"
    "24例；GSE33630，44对），采用limma拟合配对模型，差异标准为Benjamini–Hochberg校正P<0.05且|log2"
    "倍数变化|≥1。候选集合富集采用单侧Fisher精确检验，分别以全部可测基因和疾病相关基因为背景，并考察"
    "跨队列可重复性；GSE33630按放射暴露状态行敏感性分析。另对共有成分β-谷甾醇与人源CYP3A4（6UNI）、"
    "CaV3.2（9AYG）行探索性对接。"
)

RESULTS = (
    "ZBM、XKC和XYT的疾病交集基因分别为78、842和1,595个，GO和KEGG富集主要涉及氧化应激、钙信号、"
    "MAPK信号及甲状腺癌通路。两个队列各测得20,848个基因，分别检出666和745个差异表达基因。XYT和XKC"
    "在两个队列均显著富集：以疾病相关基因为背景时，GSE60542的比值比分别为1.80和2.08，GSE33630分别为"
    "1.59和1.80，校正P均<0.001；ZBM均未达统计学显著。GSE60542的666个差异基因中，664个在GSE33630"
    "变化方向一致，536个再次达到相同标准，其中包括83个XYT候选基因，如甲状腺分化标志物DIO1和TPO"
    "（两个队列均显著下调）以及LRRK2、DPP4和FN1（上调）。按放射暴露分层后，XYT和XKC的富集仍然显著"
    "（校正P均<0.001）。POLB、CACNA1H和CYP3A4均未达到联合差异表达标准。β-谷甾醇与CYP3A4和CaV3.2的"
    "最低对接评分分别为−9.12和−9.68 kcal/mol，作探索性结果解读。"
)

CONCLUSION = (
    "XYT和XKC候选基因在两个独立配对甲状腺乳头状癌队列的差异表达基因中均显著富集，以疾病相关基因为"
    "背景后结论不变。83个XYT候选基因（包括甲状腺分化标志物DIO1和TPO）在两个队列均达到差异表达标准，"
    "可作为后续功能研究优先考察的疾病相关基因。ZBM的富集证据尚不明确，POLB、CACNA1H和CYP3A4均未达到"
    "联合差异表达标准。上述结果支持候选基因集与PTC表达特征之间的疾病关联，而非草药的治疗效应；所选"
    "中药成分对甲状腺癌的作用仍需在成分明确的制剂和接受处理的模型中验证。"
)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

doc.add_heading("消瘿汤核心药组在甲状腺癌中的网络药理学及独立转录组验证", level=1)
doc.add_paragraph("结构化「材料与方法、结果、结论」投稿参考（每部分均控制在500字以内）")

for title, text in [("材料与方法", METHODS), ("结果", RESULTS), ("结论", CONCLUSION)]:
    doc.add_heading(title, level=2)
    doc.add_paragraph(text)
    doc.add_paragraph(f"（本部分字数：{len(text)}字）")

out = "output/20261027_27_before_final_review/07_结构化投稿内容_参考.docx"
doc.save(out)
print("saved:", out)
for title, text in [("材料与方法", METHODS), ("结果", RESULTS), ("结论", CONCLUSION)]:
    print(title, len(text), "字")
