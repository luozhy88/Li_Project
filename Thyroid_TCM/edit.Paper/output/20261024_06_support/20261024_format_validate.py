from pathlib import Path
import re,json,zipfile,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
out=Path(__file__).resolve().parent.parent
for name,eng in [('20261024_04_article_reviewed.docx',True),('20261024_05_复审与修改记录.docx',False)]:
 d=Document(out/name)
 for sec in d.sections:
  sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
  sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.85)
 for key in ['Normal','Body Text','First Paragraph']:
  if key not in d.styles:continue
  st=d.styles[key];st.font.name='Times New Roman' if eng else 'Noto Sans CJK SC';st.font.size=Pt(12 if eng else 10.5)
  st.paragraph_format.line_spacing=2 if eng else 1.25
  st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
 for p in d.paragraphs:
  if '[AUTHOR CONFIRM:' in p.text:
   for r in p.runs:
    if '[AUTHOR CONFIRM:' in r.text:r.font.highlight_color=WD_COLOR_INDEX.YELLOW
  if p.text.startswith('REVIEWED WORKING DRAFT'):
   for r in p.runs:r.font.color.rgb=RGBColor.from_string('555555')
  if eng and p.style.name.startswith('Heading'):
   p.paragraph_format.keep_with_next=True
   if p.text in ['Abstract','Introduction','References']:p.paragraph_format.page_break_before=True
  if p._p.xpath('.//w:drawing'):p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=False
  if eng and p.text.startswith('Figure '):p.paragraph_format.line_spacing=1.15
 for shape in d.inline_shapes:
  scale=min(1,Inches(6.5)/shape.width,Inches(8.4)/shape.height)
  shape.width=int(shape.width*scale);shape.height=int(shape.height*scale)
 foot=d.sections[0].footer.paragraphs[0];foot.alignment=2
 if not foot._p.xpath('.//w:fldSimple'):
  f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
 d.save(out/name)
 with zipfile.ZipFile(out/name) as z:assert z.testzip() is None
s=(out/'20261024_04_article_reviewed.md').read_text()
for term in ['unreconciled','prescription mismatch','Prescription inconsistency','provisionally','provisional dataset','Prescription reconciliation','repeat affected analyses']:
 assert term not in s,term
assert len(Document(out/'20261024_04_article_reviewed.docx').inline_shapes)==3
main,refs=s.split('# References');nums=[]
for token in re.findall(r'\^([\d,–]+)\^',main.split('# Introduction')[1]):
 for term in token.split(','):
  if '–' in term:
   a,b=map(int,term.split('–'));nums.extend(range(a,b+1))
  else:nums.append(int(term))
assert list(dict.fromkeys(nums))==list(range(1,21)),nums
assert len(re.findall(r'^\d+\. ',refs,re.M))==20
assert refs.count('doi:')==20
root=out.parents[1]
for item in json.loads((out/'20261024_03_support/20261024_source_manifest.json').read_text()):
 assert hashlib.sha256((root/item['path']).read_bytes()).hexdigest()==item['sha256']
checks={'date':'2026-10-24','scope':'Author-confirmed formula-derived core-herb selection; no formula mismatch objection retained.','docx_packages':'PASS','figures':3,'references_in_order':20,'source_hashes':'unchanged','new_docking_or_enrichment':'not performed','abstract_words':len(s.split('# Abstract')[1].split('**Keywords')[0].split()),'main_words_including_captions_queries':len(s.split('# Introduction')[1].split('# Conflict')[0].split())}
(out/'20261024_06_support/20261024_validation.json').write_text(json.dumps(checks,indent=2))
print(checks)
