from pathlib import Path
import re,subprocess,shutil
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml.ns import qn
R=Path('/home/data1/zhiyu/data/Github/Li_Project/Thyroid_TCM/edit.Paper/output');A=R/'20261025_25_independent_cohort';P=R/'20261025_19_JCPSP_submission'
# Correct two first-cohort-only discussion statements now that the second cohort is included.
fixes={
'Although ZBM enrichment was not statistically significant, it included only 77 measured genes and four differentially expressed genes, giving a wide confidence interval.':'ZBM included only 77 measured genes, with four differentially expressed genes in GSE60542 and two in GSE33630; its enrichment estimates had wide confidence intervals.',
'POLB showed a statistically significant change of less than 0.1 log2 units, whereas CACNA1H and CYP3A4 showed no significant expression difference.':'POLB showed small but statistically significant decreases (log2FC =−0.070 in GSE60542 and −0.116 in GSE33630), whereas CACNA1H and CYP3A4 showed no significant expression difference.'}
enpath=A/'19_manuscript_two_cohorts.md';en=enpath.read_text()
for old,new in fixes.items():assert old in en;en=en.replace(old,new)
enpath.write_text(en)
(R/'20261025_24_投稿整理归档/10_editable_sources/01_Manuscript_JCPSP_preparation.md').write_text(en)
d=Document(P/'01_Manuscript_JCPSP_preparation.docx')
for para in d.paragraphs:
 for old,new in fixes.items():
  if old in para.text:para.text=para.text.replace(old,new)
d.save(P/'01_Manuscript_JCPSP_preparation.docx')
z=(R/'20261025_22_article_中文版本合并核对.md').read_text().split('# 参考文献')[0]
z=z.replace('小瘿汤','消瘿汤')
a=z.index('# 摘要');b=z.index('# 引言')
abstract='''# 摘要

**目的：**分析消瘿汤核心药组（XYT）的关联基因，并在两个独立的甲状腺乳头状癌（PTC）队列中评估其表达特征。

**研究设计：**基于公共数据库及配对肿瘤–正常组织表达数据的二次分析。

**方法：**将XYT、浙贝母（ZBM）和夏枯草（XKC）的关联基因与GeneCards疾病基因列表取交集，开展功能富集分析。分别采用配对limma模型分析GSE60542的24例患者和GSE33630的44例患者。差异表达标准为Benjamini–Hochberg校正P<0.05且|log2倍数变化|≥1。分别以全部可测基因和其中的疾病相关基因为背景，检验候选集合的富集情况。

**结果：**ZBM、XKC和XYT的疾病交集分别包含78、842和1,595个基因。两个队列各测得20,848个基因，其中分别有666和745个达到差异表达标准。XYT和XKC在两个队列中均显著富集。采用疾病限制背景时，GSE60542的比值比分别为1.80和2.08，GSE33630分别为1.59和1.80，校正P均<0.001。初始666个差异表达基因中，664个在新队列的变化方向一致，536个再次达到相同标准。ZBM富集未达统计学显著性。POLB、CACNA1H和CYP3A4在两个队列中均未达到联合差异表达标准。

**结论：**XYT和XKC关联基因在PTC组织差异表达基因中的富集得到重复验证，为进一步开展功能研究提供了候选基因。

**关键词：**甲状腺肿瘤；中医药；基因本体；计算生物学；转录组

'''
z=z[:a]+abstract+z[b:]
z=z.replace('再在配对PTC与正常甲状腺样本中考察预先确定的候选集合','再在两个独立队列的配对PTC与正常甲状腺样本中考察预先确定的候选集合')
old=next(p for p in z.split('\n\n') if p.startswith('外部表达分析采用'))
new='''表达分析使用两个Affymetrix GPL570队列：来自比利时和法国的GSE60542，^10^以及来自乌克兰切尔诺贝利组织库的GSE33630。^NEW^分别分析提交者提供的fRMA和RMA标准化log2表达矩阵。仅纳入配对的原发PTC与非肿瘤甲状腺组织。GSE60542中同一患者同类组织的重复组织块取平均。GSE33630根据样本名称建立配对，并用临床组织库编号核实；排除5个不配对肿瘤、1个不配对正常样本及11个未分化癌。两个队列没有重复的GEO样本标识符。排除基因注释缺失或不明确的探针，同一基因的多个探针在log2尺度取平均，每个队列均获得20,848个基因。使用R 4.3.1、limma 3.58.1及稳健经验贝叶斯方法，拟合配对模型`~ patient_id + tissue_group`。^11^正log2倍数变化表示肿瘤中表达升高。差异标准为Benjamini–Hochberg校正P<0.05且|log2FC|≥1。沿用原XYT、XKC和ZBM集合，采用单侧Fisher精确检验，分别以全部可测基因和其中10,655个疾病相关基因为背景。每个队列内的6项检验统一校正P值。条件比值比同时报告双侧精确95%置信区间。通过Spearman相关、变化方向及再次达到相同差异标准的情况评估跨队列一致性。GSE33630按放射暴露状态开展敏感性分析，采用相同配对模型和阈值，12项富集检验统一校正P值。'''
z=z.replace(old,new)
# Move existing supplementary images to an appendix, matching the current English numbering.
append=[]
for old,n in [(1,1),(5,2)]:
 pattern=r'!\[\]\([^\n]*Figure'+str(old)+r'\.png\)[^\n]*\n\n\*\*图'+str(old)+r'：\*\*[^\n]+'
 match=re.search(pattern,z);assert match,old
 append.append(match.group().replace('**图'+str(old)+'：**','**补充图S'+str(n)+'：**'))
 z=z[:match.start()]+z[match.end():]
mp={1:'补充图S1',2:'图1',3:'图2',4:'图3',5:'补充图S2'}
z=re.sub(r'图([1-5])',lambda m:mp[int(m[1])],z)
newresult='''GSE33630纳入44对匹配组织，共发现745个差异表达基因，其中407个在肿瘤中上调、338个下调。XYT的1,518个可测候选基因中有112个差异基因，XKC为68/805。以全部可测基因为背景时，比值比分别为2.35和2.64；以疾病相关可测基因为背景时，分别为1.59（95%CI 1.27–1.97）和1.80（1.36–2.35），后两项检验的校正P均为4.99 × 10^-5^。ZBM包含2/77个差异基因，未显著富集，两种背景的校正P均为0.913。20,848个基因在两个队列中的log2倍数变化相关，Spearman ρ=0.786。初始666个差异基因中，664个方向一致，536个在GSE33630再次达到相同标准，其中包括83/111个XYT基因和51/69个XKC基因（补充图S3）。POLB、CACNA1H和CYP3A4仍未达到联合标准。放射暴露及未暴露两层各22对样本，XYT和XKC在两层中均保持显著富集，校正P均<0.001，ZBM仍不显著。'''
anchor=next(p for p in z.split('\n\n') if p.startswith('β-谷甾醇'))
z=z.replace(anchor,newresult+'\n\n'+anchor,1)
z=z.replace('在GSE60542中，XYT和XKC候选基因','在两个队列中，XYT和XKC候选基因')
z=z.replace('ZBM富集未达统计学显著性，但仅有77个可测基因及4个差异基因，置信区间较宽。','ZBM仅有77个可测基因，在GSE60542和GSE33630中分别有4个和2个差异基因，富集估计的置信区间较宽。')
z=z.replace('POLB虽存在统计学差异，但变化小于0.1个log2单位；','POLB在两个队列中均呈较小但有统计学意义的下降，GSE60542和GSE33630的log2FC分别为−0.070和−0.116；')
z=z.replace('表达队列仅代表PTC，组织成分差异也可能参与所观察到的变化。','两个队列均代表PTC，组织成分差异可能影响表达变化。GSE33630患者较年轻，年龄中位数16.82岁，范围12.38–24.55岁，放射暴露与未暴露者人数相等。跨队列一致性支持疾病关联的可重复性，但不能证明对草药治疗的响应。')
z=z.replace('XYT和XKC候选基因在配对PTC组织差异表达基因中富集，采用疾病限制背景后结论仍成立。','XYT和XKC候选基因在两个独立队列的配对PTC组织差异表达基因中均富集，采用疾病限制背景后结论仍成立。')
z=z.replace('GEO编号GSE60542及GPL570获取。','GEO编号GSE60542、GSE33630及GPL570获取。跨队列比较及放射暴露分层结果也见补充材料。')
z=z.replace('形式分析','统计分析').replace('写作——审阅与编辑','稿件审阅与修改')
z=re.sub(r'\^(\d+(?:,\d+)*)\^',lambda m:'^'+','.join(str(int(x)+1 if int(x)>=11 else int(x)) for x in m[1].split(','))+'^',z).replace('^NEW^','^11^')
z=z.replace('# 摘要','中文审阅说明：本稿与当前英文投稿主稿同步，包含GSE60542和GSE33630两个队列。供作者理解与审阅，不作为英文期刊投稿附件。作者姓名、单位英文信息及参考文献保留英文；补充图集中附于文后，图号与英文稿一致。\n\n# 摘要',1)
# Use identical reference entries from the current English manuscript.
z+='\n# 参考文献\n'+en.split('# References\n',1)[1]
z+='\n# 补充图（供对照阅读）\n\n'+'\n\n'.join(append)
z+='\n\n![]('+str(P/'09_figures/06_Supplementary_Figure_S3.png')+')\n\n**补充图S3：**GSE33630独立重复验证，共44对匹配组织。（A）配对limma差异表达结果，标准为BH校正P<0.05且|log2FC|≥1。（B）以全部20,848个可测基因或其中10,655个疾病相关基因为背景的候选集合富集。点为条件比值比，误差线为双侧精确95%置信区间，q为该队列6项单侧Fisher检验经BH校正后的P值。（C）两个队列的基因log2倍数变化；蓝点表示GSE60542中符合差异标准的666个基因，其中664个在GSE33630方向一致，536个方向一致且再次达到相同标准。相关系数依据全部20,848个基因计算。\n'
source=A/'25_中文完整审阅稿_双队列.md';source.write_text(z)
out=P/'07_中文完整审阅稿_双队列.docx'
subprocess.run(['/home/zhiyu/miniconda3/envs/test/bin/pandoc',str(source),'--from','markdown+superscript','--resource-path',str(R),'-o',str(out)],check=True)
d=Document(out)
for sec in d.sections:
 sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.85)
for name in ['Normal','Body Text','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Times New Roman';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
 if name in ['Normal','Body Text']:st.font.size=Pt(10.5);st.paragraph_format.line_spacing=1.25
for para in d.paragraphs:
 if para.text in ['摘要','补充图（供对照阅读）']:para.paragraph_format.page_break_before=True
for shape in d.inline_shapes:
 scale=min(1,Inches(6.4)/shape.width,Inches(7.3)/shape.height);shape.width=int(shape.width*scale);shape.height=int(shape.height*scale)
d.save(out)
text='\n'.join(x.text for x in d.paragraphs)
assert all(x in text for x in ['GSE33630','44','745','664','536','−0.116','Xing Zhang','21MC1930500','ZX0551-25','2025JZ02','补充图S3'])
assert not any(x in text for x in ['Pengfei Yang','王文春','小瘿汤','2024_09_05'])
assert len(d.inline_shapes)==6
assert len(re.findall(r'^\d+\.',z.split('# 参考文献')[1].split('# 补充图')[0],re.M))==25
shutil.make_archive(str(P),'zip',P.parent,P.name)
print('Chinese complete manuscript created and verified; 6 figures, 25 references; ZIP updated.')
