from pathlib import Path
from zipfile import ZipFile
import io, re
import fitz
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

base = Path(__file__).parent
work = base / '_combined_materials_tmp'
work.mkdir(exist_ok=True)
with ZipFile(base / '05_Supplementary_tables.zip') as z:
    z.extractall(work)

xlsx = base / '10_Supplementary_tables_all.xlsx'
with pd.ExcelWriter(xlsx, engine='openpyxl') as writer:
    readme = (work / '00_README.txt').read_text(errors='replace')
    pd.DataFrame({'README': readme.splitlines()}).to_excel(writer, sheet_name='README', index=False)
    for csv in sorted(work.glob('*.csv')):
        name = re.sub(r'[^A-Za-z0-9_]', '_', csv.stem)[:31]
        try:
            df = pd.read_csv(csv, low_memory=False)
        except Exception:
            df = pd.read_csv(csv, low_memory=False, encoding='latin1')
        df.to_excel(writer, sheet_name=name, index=False)
wb = load_workbook(xlsx)
for ws in wb.worksheets:
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
    for col in ws.iter_cols():
        n = max(len(str(c.value)) if c.value is not None else 0 for c in col[:100])
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(42, max(10, n + 2))
wb.save(xlsx)

fig = base / '09_figures'
items = [
    ('Figure 1', fig/'01_Figure_1.png'), ('Figure 2', fig/'02_Figure_2.png'),
    ('Figure 3', fig/'03_Figure_3.png'), ('Supplementary Figure S1', fig/'04_Supplementary_Figure_S1.png'),
    ('Supplementary Figure S2', fig/'05_Supplementary_Figure_S2.png'), ('Supplementary Figure S3', fig/'06_Supplementary_Figure_S3.png'),
    ('GSE60542 QC PCA', work/'20_QC_PCA.pdf'), ('GSE60542 volcano', work/'21_volcano.pdf'),
    ('GSE60542 enrichment', work/'24_transcriptome_figure.png'), ('GSE33630 QC PCA', work/'34_GSE33630_PCA.pdf')]
W, H = 1654, 2339
try: font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 34)
except Exception: font = ImageFont.load_default()
pages = []
for title, path in items:
    if path.suffix.lower() == '.pdf':
        d = fitz.open(path)
        imgs = [Image.open(io.BytesIO(p.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False).tobytes('png'))).convert('RGB') for p in d]
    else: imgs = [Image.open(path).convert('RGB')]
    for i, image in enumerate(imgs, 1):
        page = Image.new('RGB', (W, H), 'white')
        draw = ImageDraw.Draw(page)
        draw.text((70, 45), title if len(imgs) == 1 else f'{title} — page {i}', fill='black', font=font)
        scale = min((W-140)/image.width, (H-160)/image.height)
        image = image.resize((int(image.width*scale), int(image.height*scale)), Image.Resampling.LANCZOS)
        page.paste(image, ((W-image.width)//2, 115+(H-160-image.height)//2))
        pages.append(page)
pdf = base / '11_Supplementary_figures_all.pdf'
pages[0].save(pdf, 'PDF', resolution=150.0, save_all=True, append_images=pages[1:])
print('XLSX', xlsx, xlsx.stat().st_size)
print('PDF', pdf, pdf.stat().st_size, 'pages', len(pages))
